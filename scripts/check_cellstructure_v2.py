#!/usr/bin/env python3
"""Meaningful rebuild checks: PDF text retention and discovery order, separate from visual review."""
import json,re,subprocess,hashlib
from pathlib import Path
from pptx import Presentation
from build_cellstructure_v2 import NAMES,ROOT
reports=[]
for n,name in NAMES.items():
 p=ROOT/'outputs/drafts/v2'/f'{name}.pptx';pdf=ROOT/'qa/renders/cellstructure_v2'/f'{name}.pdf';deck=Presentation(p)
 pages=subprocess.check_output(['pdftotext','-layout',str(pdf),'-'],text=True).split('\f');errors=[]
 def words(t):return set(re.findall(r'[a-z]{4,}',t.lower()))
 for i,s in enumerate(deck.slides):
  texts=[]
  for para in s.notes_slide.notes_text_frame.paragraphs:
   if para.font.name!='Calibri' or any(r.font.name!='Calibri' for r in para.runs):errors.append({'slide':i+1,'notes_font':'Expected explicit Calibri'})
  for sh in s.shapes:
   if sh.has_text_frame:texts.append(sh.text)
   if sh.has_table:texts.extend(c.text for row in sh.table.rows for c in row.cells)
  missing=words('\n'.join(texts))-words(pages[i])
  if missing:errors.append({'slide':i+1,'missing_pdf_words':sorted(missing)})
 if n==6:
  rec=next(i for i,s in enumerate(deck.slides,1) if any(sh.has_text_frame and sh.text=='Record YOUR view' for sh in s.shapes))
  comp=next(i for i,s in enumerate(deck.slides,1) if any(sh.has_text_frame and sh.text=='Now compare with a drawing' for sh in s.shapes))
  if comp<=rec:errors.append('Practical comparison precedes pupil recording')
  if f'slide {comp}' not in deck.slides[0].notes_slide.notes_text_frame.text:errors.append('Comparison warning has a stale slide number')
 if n==4:
  for term in ['Cell wall','Cell membrane','Nucleus','Chloroplast','Large central vacuole','Mitochondrion','Rough ER','Ribosomes','Golgi apparatus']:
   if not any(sh.has_text_frame and sh.text==term+' · the job' for s in deck.slides for sh in s.shapes):errors.append('Missing function reveal: '+term)
 reports.append({'lesson':n,'deck':name,'slides':len(deck.slides),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'pdf_text_and_sequence':'PASS' if not errors else 'FAIL','errors':errors})
report={'lessons':reports,'limits':'Text extraction catches missing words, not every overlap, wrong label endpoint, scientific error or PowerPoint difference. Visual and teacher review are separate.'}
(ROOT/'qa/CellStructure_v2_content_checks.json').write_text(json.dumps(report,indent=2));print(json.dumps(report,indent=2));raise SystemExit(any(x['errors'] for x in reports))
