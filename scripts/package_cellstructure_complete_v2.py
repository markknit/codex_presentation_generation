#!/usr/bin/env python3
"""Package twelve review drafts and source-aligned print tasks. No release approval is implied."""
import argparse,copy,hashlib,json,shutil,subprocess,zipfile
from pathlib import Path
from pptx import Presentation
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.table import WD_ROW_HEIGHT_RULE
from build_cellstructure_remaining_v2 import ROOT,ALL_NAMES,REMAINING_NAMES,PREP,SCI
from cell_specialisation_visuals import FEATURES,FUNCTIONS

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def doc(title):
 d=Document();d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(10)
 d.sections[0].header.paragraphs[0].text='SeaVuria · ARES Education · Grade 10 Biology · REVIEW DRAFT'
 d.sections[0].footer.paragraphs[0].text='© 2026 SeaVuria & ARES Education · CC BY-NC 4.0 · Adapted 24 October 2026'
 d.add_heading(title,0);return d
def table(d,headers,rows,height=None):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for row in rows:
  rr=t.add_row()
  for c,v in zip(rr.cells,row):c.text=str(v)
  if height:rr.height=Inches(height);rr.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
 return t
def blank(d,title,height=3):
 d.add_heading(title,1);t=d.add_table(rows=1,cols=1);t.style='Table Grid';t.rows[0].height=Inches(height);t.rows[0].height_rule=WD_ROW_HEIGHT_RULE.EXACTLY

def main():
 p=argparse.ArgumentParser();p.add_argument('--out',required=True,type=Path);a=p.parse_args();a.out.mkdir(exist_ok=True,parents=True)
 out=a.out;work=ROOT/'.cache/package_complete_v2';work.mkdir(exist_ok=True)
 def save(d,name):d.save(out/name)
 def render(n):return ROOT/'qa/renders'/('cellstructure_complete_v2' if n in REMAINING_NAMES else 'cellstructure_v2')
 def extract(n,pages,name):
  parts=[]
  for k in pages:
   dest=work/f'l{n}-{k}.pdf';subprocess.run(['pdfseparate','-f',str(k),'-l',str(k),str(out/f'{ALL_NAMES[n]}.pdf'),str(dest)],check=True);parts.append(str(dest))
  subprocess.run(['pdfunite',*parts,str(out/name)],check=True)
 for n,name in sorted(ALL_NAMES.items()):
  src=ROOT/'outputs/drafts/v2'/f'{name}.pptx';report=json.loads((ROOT/'qa'/f'{name}_validation.json').read_text())
  assert report['structural_status']=='PASS' and report['sha256']==sha(src)
  for f in [src,render(n)/f'{name}.pdf']:shutil.copy2(f,out/f.name)
  d=doc(f'Lesson {n} · teacher notes');d.add_paragraph('Teacher-only. Slides 2–3 are hidden in Slide Show but included here and in full PDFs. Calibri is specified; Linux previews use Carlito. Native PowerPoint/Presenter View and teacher approval remain open.')
  if n in PREP:d.add_heading('Preparation',1);d.add_paragraph(PREP[n]);d.add_heading('Science and source decisions',1);d.add_paragraph(SCI[n])
  for i,s in enumerate(Presentation(src).slides,1):
   d.add_page_break();d.add_heading(f'Slide {i}'+(' · HIDDEN' if s._element.get('show')=='0' else ''),1)
   d.add_picture(str(render(n)/f'l{n}/gpu/slide-{i:02d}_preview.png'),width=Inches(6.2))
   for part in s.notes_slide.notes_text_frame.text.split('\n\n'):
    if part.strip():d.add_paragraph(part.strip())
  save(d,f'CellStructure_L{n}_v2_TeacherNotes.docx')
 # Print the diagrams at the correct lesson phase; answer-bearing resources stay teacher-controlled.
 references={2:list(range(10,17)),3:[8,9,10],4:[8],5:[8,9,10,11],7:[11],8:list(range(8,17)),9:list(range(8,15)),10:[8,9,10],11:[8]}
 for n,pages in references.items():extract(n,pages,f'CellStructure_L{n}_v2_ReferenceSheets.pdf')
 quick=[]
 for n,pages in {3:[8,10],5:[8,9],7:[10,14],8:[10,13,22],9:[8,11,13],10:[8,9],11:[8,10],12:[8,14,19,21]}.items():
  for k in pages:
   dest=work/f'quick-{n}-{k}.pdf';subprocess.run(['pdfseparate','-f',str(k),'-l',str(k),str(out/f'{ALL_NAMES[n]}.pdf'),str(dest)],check=True);quick.append(str(dest))
 subprocess.run(['pdfunite',*quick,str(out/'CellStructure_Remaining_v2_QuickReview.pdf')],check=True)
 # Retain the reviewed original L4/L6 blank tasks and design guide from the previous snapshot.
 with zipfile.ZipFile(ROOT/'outputs/review/CellStructure_Lessons1-2-4-6_v2_Review.zip') as z:
  for f in ['CellStructure_L4_v2_BlankTasks.docx','CellStructure_L6_v2_ObservationSheet.docx','CellStructure_v2_DesignGuide.docx']:(out/f).write_bytes(z.read(f))
 d=doc('Lesson 3 · compare two cell models');d.add_paragraph('Name: __________________ Date: __________\nComplete individually before partner comparison. Record typical examples, not universal rules.')
 table(d,['Feature','E. coli example','Body-cell example'],[[k,'',''] for k in ['DNA / nuclear envelope','Membrane-bound organelles','Ribosomes','Cell wall / material','Typical size']],.65);blank(d,'Three sentences · structural differences and Amina’s story',1.4);save(d,'CellStructure_L3_v2_BlankTasks.docx')
 d=doc('Lesson 5 · animal-cell prediction');d.add_paragraph('Draw and label your prediction before labelled references are shown. Star uncertain structures.');blank(d,'Unlabelled outline · draw within this boundary',4.5)
 d.add_page_break();d.add_heading('Compare typical examples',1);table(d,['Structure + function','Plant model','Animal model'],[[k,'',''] for k in ['Nucleus, mitochondria, ribosomes','ER, Golgi, membrane','Lysosomes / lytic vacuole','Centrioles','Wall, chloroplast, central vacuole']],.65)
 d.add_heading('Peer feedback',1);table(d,['What Went Well · specific science','Even Better If · specific correction'],[['','']],1);save(d,'CellStructure_L5_v2_BlankTasks.docx')
 d=doc('Lesson 7 · my own animal-cell evidence');d.add_paragraph('Name: __________ Date: ______\nUse the school-approved prepared slide. No swabbing or staining in this route. Draw what YOU can identify; uncertainty is a valid record.')
 d.add_paragraph('Specimen / slide label: __________\nEyepiece ____ × Objective ____ × Total optical magnification ____ ×');blank(d,'Large scientific drawing · pencil, ruled labels, no shading',4.2)
 d.add_page_break();d.add_heading('Compare my actual records',1);table(d,['Record','Onion (L6)','Prepared animal cell (L7)'],[[k,'',''] for k in ['Identified structures','Not seen / unclear','Magnification','Evidence supporting comparison']],.75)
 blank(d,'Three-sentence evidence-based explanation',1.8);d.add_paragraph('What my observation supports: __________________\nWhat remains uncertain: __________________\nWhat comes from a reference model rather than my view: __________________');save(d,'CellStructure_L7_v2_ObservationSheet.docx')
 animals=['red blood cell','neutrophil','neuron','sperm cell','skeletal muscle fibre']
 d=doc('Lesson 8 · observation and explanation organiser');d.add_paragraph('Name: __________ Group: ____\nObserve first. Distinguish a model feature from something actually resolved in a published image.')
 table(d,['Cell','Features recorded','Absent / unclear','Questions'],[[k,'','',''] for k in animals],.8)
 d.add_page_break();d.add_heading('Explain after observation',1);table(d,['Cell','Feature','Why it helps','Cut-finger link','Sickle-cell link'],[[k,'','','',''] for k in animals],.9);d.add_paragraph('A link can be “not involved / not applicable” when justified. Sperm does not repair the cut.');save(d,'CellStructure_L8_v2_BlankOrganiser.docx')
 d=doc('Lesson 8 · background data card');d.add_paragraph('Use during Explain, after recording your observations. This is published explanatory background; it is not a result measured by our class.')
 d.add_heading('Red cells and sickling',1);d.add_paragraph('Typical adult red cells are flexible biconcave discs containing haemoglobin. They deliver oxygen to respiring tissues. Their typical lifespan is about 120 days. In sickle-cell disease, abnormal haemoglobin can form polymers; affected cells may become less deformable, obstruct small vessels and break down prematurely, contributing to anaemia. The blood-smear photograph shows shapes, not stiffness, blood flow or a diagnosis.')
 d.add_paragraph('Sources: https://www.nhlbi.nih.gov/health/sickle-cell-disease ; https://openstax.org/books/anatomy-and-physiology-2e/pages/18-3-erythrocytes')
 d.add_heading('Other specialised cells',1)
 for kind in animals[1:]:d.add_paragraph(kind.capitalize()+': '+FUNCTIONS[kind])
 d.add_paragraph('Neutrophil lifespan varies with context; no fixed six-hour rule or unverified circulation count is used. The published neural image is fluorescent rat hippocampal tissue, a different method and specimen from school microscopy. Do not interpret staining markers from this card.')
 d.add_heading('Image credits',1)
 d.add_paragraph('OpenStax Biology for AP Courses Fig. 3.29 (modified Ed Uthman; scale Matt Russell) and Biology 2e Fig. 35.6 (modified Dr Maryam Faiz et al., University of Barcelona), CC BY 4.0. Exact URLs and hashes: published_evidence.json.');save(d,'CellStructure_L8_v2_DataCard.docx')
 plants=['root hair cell','guard cells','xylem vessel','phloem sieve tube','palisade cell']
 d=doc('Lesson 9 · observation table');table(d,['Cell','Structural features shown','Function hypothesis'],[[k,'',''] for k in plants],.9);d.add_paragraph('Guard-cell state images are hypothetical models, not photographs or measured data from Turkana/Kericho.');blank(d,'Which model surprised me? Why?',1);save(d,'CellStructure_L9_v2_BlankTasks.docx')
 d=doc('Lesson 9 · function and location cards');d.add_paragraph('Teacher: cut the five diagram pages (L9 ReferenceSheets pages 1–5) into five diagram cards, then cut the ten cards below. Shuffle the three sets separately: 15 cards per group. Diagram cards are labelled as in the V2 source; function/location matches still require structural reasoning. Retain this instruction page.')
 d.add_page_break();table(d,['FUNCTION CARDS · cut each row'],[[FUNCTIONS[k]] for k in plants],.9)
 d.add_page_break();table(d,['LOCATION CARDS · cut each row'],[[v] for v in ['Root epidermis behind the growing tip, in contact with soil.','Leaf epidermis, paired around a stoma.','Water-conducting vascular tissue through root, stem and leaf.','Sugar-conducting vascular tissue between sources and sinks.','Photosynthetic mesophyll, commonly near the upper surface of a dicot leaf.']],.9);save(d,'CellStructure_L9_v2_MatchingCards.docx')
 d=doc('Lesson 10 · two-sided evidence card');d.add_paragraph('Side A · Use L10 ReferenceSheets page 1 for the drawn biological pathway. These are schematic models, not micrographs.')
 d.add_heading('School analogy · a comparison with limits',1);table(d,['Example','Contribution in this analogy'],[['One pupil','Individual contribution'],['Class of 42 pupils','A group learning together'],['Mwangaza Secondary School','Different classes and roles coordinated'],['Nairobi County Education','Several schools coordinated'],['Kenyan society','Many systems working together']],.55)
 d.add_paragraph('Source classroom example, not a real survey. A head teacher does not model cell signalling exactly. Some body cells migrate; a tissue can include several cell types and extracellular material. An analogy organises ideas; it does not prove a biological mechanism.')
 d.add_page_break();d.add_heading('Side B · healing-cell profiles',1)
 profiles={'platelet':'Small activated fragments change shape and aggregate, forming a plug and supporting clotting. Fibrin stabilises the clot.','neutrophil':'Flexible cell with a lobed nucleus and digestive compartments; engulfs microbes and digests engulfed material.','macrophage':'Irregular flexible membrane and digestive compartments support engulfment of microbes, dead cells and debris.','fibroblast':'Protein-synthesis and secretion machinery, including rough ER and Golgi, supports collagen production in connective tissue.','epidermal cells':'Living cells can divide and migrate to restore the surface. Their cytoskeleton and lack of a rigid wall permit changes of shape.'}
 for k,v in profiles.items():d.add_heading(k.capitalize(),2);d.add_paragraph(v)
 save(d,'CellStructure_L10_v2_EvidenceCard.docx')
 d=doc('Lesson 10 · organisation table');table(d,['Level','Biological example','School analogy','Function'],[[k,'','',''] for k in ['Cell','Tissue','Organ','Organ system','Organism']],.9);blank(d,'Five sentences · why one cell cannot heal the cut alone',2);save(d,'CellStructure_L10_v2_BlankTasks.docx')
 d=doc('Lesson 11 · poster plan and feedback');d.add_paragraph('Use A2 paper for the final poster. Plan from your own prior models and evidence.');table(d,['A · PLANT CELL','B · ANIMAL CELL'],[['8+ structures, function notes\n\n\n\n','6+ structures, comparison\n\n\n\n']],2);table(d,['C · THREE HEALING CELLS / FRAGMENTS','D · ORGANISATION'],[['Feature → job → cut\n\n\n\n','Cell → tissue → organ → system → organism\n\n\n\n']],2)
 d.add_heading('Scientific feedback',1);d.add_paragraph('Glow · a specific accurate link: __________________\nGrow · a scientific gap or unclear label: __________________\nRevision and evidence: __________________');save(d,'CellStructure_L11_v2_PosterPlan.docx')
 assessment(out,save)
 # QA and logs travel with the review package; all working drafts remain ignored in the repository.
 evidence=out/'qa';evidence.mkdir(exist_ok=True)
 for f in (ROOT/'qa').glob('*.json'):
  if f.name.startswith('CellStructure_') and not f.name.endswith('_package.json') and ('v2' in f.name or 'inventory' in f.name):shutil.copy2(f,evidence/f.name)
 for name in ALL_NAMES.values():
  for suffix in ['manifest','validation','gpu']:
   f=ROOT/'qa'/f'{name}_{suffix}.json';assert f.exists(),f;shutil.copy2(f,evidence/f.name)
 logs=out/'build_logs';logs.mkdir(exist_ok=True)
 for f in (ROOT/'build_logs').glob('*_v2.md'):shutil.copy2(f,logs/f.name)
 for f in ['provenance.json','published_evidence.json']:shutil.copy2(ROOT/'assets/cellstructure_v2'/f,out/f)
 shutil.copy2(ROOT/'content/cellstructure_v2/DESIGN.md',out/'DESIGN.md')
 (out/'START_HERE.txt').write_text('CELL STRUCTURE · COMPLETE VERSION 2 · LESSONS 1–12\nREVIEW DRAFTS · 276 slides · 840 scheduled minutes\n\nNew/rebuilt: L3, L5, L7–12 (166 slides). L1/L2/L4/L6 retain the previous reviewed draft versions. Start with the 20-slide Remaining QuickReview, then inspect each separate PPTX/full PDF and teacher notes. PDFs include hidden teacher slides 2–3. Do not project teacher notes or the assessment marking guide.\n\nUse print sheets only at the stated phase. Labelled references follow Predict; L7 comparison illustration follows actual recording. L8 DataCard is for Explain. L9 has five diagram cards in ReferenceSheets pages 1–5 plus ten cut-out cards in MatchingCards. L10 EvidenceCard is two-sided; science graphics are schematic.\n\nL7 uses a school-approved prepared epithelial-cell slide, with no learner swabs/stain. L9 needs a seedling prepared 7–10 days ahead if available. L12 begins the five-section assessment with 25 minutes for Sections 1–3; schedule supervised follow-up for 4–5. The assessment preserves source sections/rubric with logged scientific and route corrections.\n\nOpen before classroom release: teacher/source-owner review, local practical readiness, native PowerPoint/Slide Show/Presenter View, exact Calibri, and ARES link/content/offline checks. L6 chemical/PPE/calibration decisions remain open. Sources disagree on numerical sub-strand designation; Cell Structure is used pending owner resolution. GPU preview resizing/pixel metrics executed on NVIDIA GB10; CPU-only PowerPoint/vector/DOCX assembly, LibreOffice conversion, Poppler rasterisation, encoding, I/O and ZIP are disclosed. No semantic GPU or teaching-effectiveness approval is claimed.\n\nDrafts are local. The previously pushed four-lesson snapshot is preserved; no new push is implied by this package.\n')
 manifest={str(f.relative_to(out)):{'bytes':f.stat().st_size,'sha256':sha(f)} for f in sorted(out.rglob('*')) if f.is_file() and f.name!='package_manifest.json'}
 (out/'package_manifest.json').write_text(json.dumps(manifest,indent=2))
 bundle=out.parent/'CellStructure_Complete_Lessons1-12_v2_Review.zip'
 with zipfile.ZipFile(bundle,'w',zipfile.ZIP_DEFLATED) as z:
  for f in sorted(out.rglob('*')):
   if f.is_file():z.write(f,f.relative_to(out))
 with zipfile.ZipFile(bundle) as z:assert z.testzip() is None
 print(json.dumps({'bundle':str(bundle),'files':len(manifest)+1,'bytes':bundle.stat().st_size,'sha256':sha(bundle)},indent=2))

def assessment(out,save):
 source=Document(ROOT/'inputs/v2/Biology/SS1.1_Cell_Structure/Biology_Cell_Structure_FinalExplanation.docx')
 prompts=[t.cell(1,0).text for t in source.tables[2:7]]
 prompts[0]=prompts[0].replace('from badly prepared nyama choma','that may be transferred from contaminated food or surfaces')
 prompts[1]='In Lessons 4–7 you compared typical plant and animal cell models and recorded onion and approved prepared animal epithelial cells. (a) Compare the highlighted structures in the photosynthetic flowering-plant and typical animal models; state an important exception to a simple “plant only / animal only” rule. (b) What differences and similarities can you justify from YOUR two practical records? Distinguish “not visible” from “absent”. (c) Explain why onion bulb epidermis may lack visible chloroplasts. (d) Explain how lacking a rigid cell wall helps animal cells that close the cut. (e) Estimate the length of an onion cell from your own calibrated field measurement and cell count, showing units and method. If a measurement was not possible, state the limitation and the method you would use. State one ultrastructural detail a suitable electron micrograph can resolve that our school light microscope cannot.'
 prompts[2]=prompts[2].replace('sickle-cell anaemia micrographs','published blood-smear image and explanatory background').replace('Does the shape of a cell affect the job it can do?','Does the shape of a cell affect the job it can do? Distinguish visible shape evidence from the mechanism supplied by the background card.')
 prompts[3]=prompts[3].replace('if all cells in one organism carry the same DNA, why do they look so different?','if most nucleated body cells share essentially the same genetic instructions, why do they look so different?').replace('a platelet or a fibroblast','an epidermal cell or a fibroblast')
 d=doc('Final Explanation · student assessment');d.add_paragraph('Name: __________________ Class: ______ Date: ______\nBegin Sections 1–3 during the 25-minute independent session. Sections 4–5 continue in a supervised follow-up under school policy. Use additional paper if needed.')
 d.add_paragraph('Driving question: Why do cells in our body look so different from each other — and how does the shape of a cell determine the job it can do? Explain Amina’s healing cut using specific lesson evidence, full sentences, correct terms and Claim–Evidence–Reasoning where requested. Use actual practical records and identify secondary evidence honestly.')
 for i,prompt in enumerate(prompts):
  d.add_page_break();d.add_heading(source.tables[i+2].cell(0,0).text,1);d.add_paragraph(prompt);blank(d,'My answer',5.1 if i!=1 else 4.3)
 d.add_page_break();d.add_heading('Source rubric',1);rubric(d,source)
 # Preserve bands and criteria, adapting only the practical-route evidence wording.
 for t in d.tables[-1:]:
  for row in t.rows:
   for c in row.cells:
    if 'onion and cheek practicals' in c.text:c.text=c.text.replace('onion and cheek practicals','onion and prepared animal-cell practicals')
 attribution(d,source);save(d,'CellStructure_L12_v2_StudentAssessment.docx')
 keys=[
 'Cell theory: organisms consist of cells; cells are the basic unit of life; cells arise from pre-existing cells. Virchow is conventionally associated with the third tenet. Surviving living epidermal cells divide and migrate; new skin does not arise from nothing. E. coli DNA occupies a nucleoid without a nuclear envelope, and there are no membrane-bound organelles; a typical nucleated body cell has these compartments. Both have ribosomes. Typical size ranges are supporting examples, not a universal diagnostic test; many E. coli strains are harmless.',
 'Compare these examples, not universal categories. A photosynthetic flowering-plant model has a wall, chloroplasts and large central vacuole; the animal model highlights lysosomes and centrioles. Plants also digest/recycle in lytic compartments, and centrioles are not absent from all plant groups. Both models have membrane, cytoplasm, ribosomes, nucleus, mitochondria, ER and Golgi. Practical answers must match the pupil’s own records; do not award invented chloroplasts, a separately resolved membrane against an onion wall, or a large vacuole merely because the teacher key suggests one. A poor view is not proof of absence. Bulb epidermal cells are not typical photosynthetic leaf cells; chloroplast absence is consistent with their specialised role, not simply a rule that every underground cell lacks chloroplasts. Animal cells can change shape and migrate with cytoskeletal support and no rigid wall. Length ≈ calibrated field diameter ÷ cell lengths spanning the same axis/field; convert mm ×1000 to μm. Judge actual measurements and uncertainty, not a fixed numerical answer. Suitable TEM can resolve internal membranes/ribosomes beyond school light-microscope resolution. The unit’s optional EM set was not supplied, so accept taught resolution limits without pretending pupils observed an EM image.',
 'Any three healing cells/fragments with a correct feature → function link: activated platelets change shape, adhere/aggregate and support clotting; neutrophils have flexible boundaries and digestive compartments for engulfment/digestion; macrophages engulf debris/microbes and digest them; fibroblast rough ER/ribosomes/Golgi support collagen synthesis/secretion; living epidermal cells divide and migrate to restore the surface. A platelet is a fragment, not a whole nucleated cell. Sickle-cell CER: claim shape/function are linked; evidence describes disc-like and crescent profiles actually visible in the published smear; reasoning uses the background that abnormal haemoglobin can polymerise, reducing deformability, obstructing small vessels and increasing premature cell breakdown. The static image alone does not measure stiffness/flow or diagnose an individual. Do not require “curved cells cannot carry oxygen”. Plant example: root-hair extension increases uptake surface; mature xylem vessel elements form a hollow lignified pathway; palisade chloroplasts support photosynthesis; sieve tubes transport source-to-sink with companion-cell support; guard-cell turgor changes the pore.',
 'Most nucleated somatic cells in one organism share essentially the same genome but express different sets/amounts of genes, producing different proteins and structures. Avoid claiming literally all DNA is identical: mature mammalian red cells and platelets lack nuclei; gametes and other biological variation are outside the simplified comparison. Valid same-organism examples are a neutrophil/fibroblast in one person or root hair/palisade in one bean plant. Epidermal cells contribute to epithelial tissue and fibroblasts to dermal connective tissue; several tissues form skin, which belongs to the integumentary system protecting the organism. Blood/immune cells cooperate through circulation and migration; they do not all become epidermis. A tissue may contain several cell types and extracellular material. No single cell supplies every repair role.',
 'Accept a coherent account of at least four healing cells/fragments linked to specific features and coordinated roles. A simplified sequence is bleeding → platelet plug/clot → overlapping immune clearance → connective-tissue repair and surface restoration. Fibrin forms from existing plasma fibrinogen through clotting reactions; ribosomes do not weave fibrin at the wound. Neutrophils/macrophages engulf and digest; fibroblasts synthesise/secrete collagen; living epidermal cells copy DNA before division and migrate. Spindle microtubules segregate chromosomes; centrioles do not directly pull DNA. Gene expression and organisation explain differing contributions. Healing stages overlap and restored skin need not be identical to the original. Reflection should refer to the learner’s authentic Lesson 1 idea and a specific evidence source, not reproduce a scripted emotional response.'
 ]
 d=doc('Final Explanation · corrected teacher marking guide');d.add_paragraph('TEACHER ONLY · source-adapted review guide, not an approved marking scheme. Keep private during assessment. Preserve the five source criteria and the 4 / 3 / 1–2 performance bands; no new point allocation or pass threshold has been invented. Judge reasoning and authentic route-specific evidence. The full assessment is not expected to finish in the first 25 minutes.')
 for i,key in enumerate(keys,1):d.add_heading(f'Section {i} · expected scientific reasoning',1);d.add_paragraph(key)
 d.add_page_break();d.add_heading('Rubric',1);rubric(d,source)
 for t in d.tables:
  for r in t.rows:
   for c in r.cells:
    if 'onion and cheek practicals' in c.text:c.text=c.text.replace('onion and cheek practicals','onion and prepared animal-cell practicals')
 d.add_heading('Adaptation record and checks',1);d.add_paragraph('Preserved five sections, source learning intent, choice of three healing cells, one plant example, same-organism comparison, four-cell final account and rubric bands. Corrected universal organelle/DNA claims, actual prepared-slide route, visibility assumptions, image-versus-mechanism reasoning and fixed example measurements. Original sources remain unchanged in inputs/v2. See per-deck build logs for source decisions.')
 d.add_paragraph('Primary science references: OpenStax Biology 2e §§4.3, 30.2; OpenStax Anatomy and Physiology 2e §§18.3, 18.5; NHLBI Sickle Cell Disease (https://www.nhlbi.nih.gov/health/sickle-cell-disease). Image provenance is supplied separately.')
 attribution(d,source);save(d,'CellStructure_L12_v2_TeacherMarkingGuide.docx')

def rubric(d,source):
 # Reflow the source table to this portrait page rather than inheriting oversized source widths.
 rows=[[c.text for c in row.cells] for row in source.tables[7].rows]
 t=table(d,rows[1],rows[2:]);t.autofit=False
 for row in t.rows:
  for j,c in enumerate(row.cells):
   c.width=Inches(1.2 if j==0 else 1.766)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(3)
    for r in p.runs:r.font.size=Pt(9);r.font.name='Calibri'
 for j,col in enumerate(t.columns):col.width=Inches(1.2 if j==0 else 1.766)


def attribution(d,source):
 d.add_page_break();d.add_heading('Attribution and licensing',1)
 for p in source.paragraphs:
  if p.text.startswith(('Curriculum Alignment:','Content & Pedagogy:','Production & Layout:','Licensing & Copyright:')):d.add_paragraph(p.text)
 d.add_paragraph('This review adaptation was assembled with Codex/OpenAI. It is independently developed; no KICD endorsement or new teacher approval is claimed. External OpenStax images retain CC BY 4.0 attribution.')

if __name__=='__main__':main()
