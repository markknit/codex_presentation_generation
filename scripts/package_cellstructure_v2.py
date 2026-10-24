#!/usr/bin/env python3
"""Package only validated v2 review drafts, not approved classroom releases."""
import argparse,json,hashlib,shutil,subprocess,zipfile
from pathlib import Path
from pptx import Presentation
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.table import WD_ROW_HEIGHT_RULE
from build_cellstructure_v2 import ROOT,NAMES,ORG
p=argparse.ArgumentParser();p.add_argument('--out',type=Path,required=True);a=p.parse_args();a.out.mkdir(parents=True,exist_ok=True)
files=[]
def keep(p):files.append(p);return p
def doc():
 d=Document();d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(10);return d
def table(d,headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for r in rows:
  for c,v in zip(t.add_row().cells,r):c.text=v
 return t
work=ROOT/'.cache/package_v2';work.mkdir(exist_ok=True)
allpdf={}
for n,name in NAMES.items():
 report=json.loads((ROOT/'qa'/f'{name}_validation.json').read_text());src=ROOT/'outputs/drafts/v2'/f'{name}.pptx'
 assert report['structural_status']=='PASS' and report['sha256']==hashlib.sha256(src.read_bytes()).hexdigest()
 for src in [src,ROOT/'qa/renders/cellstructure_v2'/f'{name}.pdf']:
  target=keep(a.out/src.name);shutil.copy2(src,target)
  if target.suffix=='.pdf':allpdf[n]=target
 # One document per lesson keeps reviewer navigation practical.
 d=doc();d.add_heading(f'Lesson {n} · teacher notes and review',0);d.add_paragraph('24 October 2026 · REVIEW DRAFT · Calibri specified; local previews use Carlito.')
 d.add_paragraph('Check native PowerPoint/Presenter View, teaching sequence, science, local timings and every build-log item before release. Slides 2–3 are hidden in Slide Show but included in PDFs and these notes.')
 if n==6:d.add_paragraph('Practical readiness: teacher must confirm stain product/concentration/SDS, PPE, waste, equipment, calibration and actual visibility. No source shortcut or chemical recipe is approved here.')
 if n==4:d.add_paragraph('Diagram-only route. Optional real EM set is not supplied. These illustrations are never evidence of an actual microscopy result.')
 for i,s in enumerate(Presentation(ROOT/'outputs/drafts/v2'/f'{name}.pptx').slides,1):
  d.add_page_break();d.add_heading(f'Slide {i}'+(' · HIDDEN' if s._element.get('show')=='0' else ''),1)
  d.add_picture(str(ROOT/f'qa/renders/cellstructure_v2/l{n}/gpu/slide-{i:02d}_preview.png'),width=Inches(6.25))
  for chunk in s.notes_slide.notes_text_frame.text.split('\n\n'):
   if chunk.strip():d.add_paragraph(chunk.strip())
 d.save(keep(a.out/f'CellStructure_L{n}_v2_TeacherNotes.docx'))
# Quick review: deliberately selected student screens, not substitute for full review PDFs.
quick=[]
for n,pages in {1:[1,7,12,18],2:[9,12,20],4:[1,8,21,25,31,34],6:[10,14,20,22]}.items():
 for page in pages:
  dest=work/f'quick-l{n}-{page:02d}.pdf';subprocess.run(['pdfseparate','-f',str(page),'-l',str(page),str(allpdf[n]),str(dest)],check=True);quick.append(str(dest))
subprocess.run(['pdfunite',*quick,str(keep(a.out/'CellStructure_v2_QuickReview.pdf'))],check=True)
# Source tasks, directly printable at the appropriate phase.
for n,pages in {2:list(range(10,17)),4:[8]}.items():
 parts=[]
 for page in pages:
  dest=work/f'print-l{n}-{page:02d}.pdf';subprocess.run(['pdfseparate','-f',str(page),'-l',str(page),str(allpdf[n]),str(dest)],check=True);parts.append(str(dest))
 subprocess.run(['pdfunite',*parts,str(keep(a.out/f'CellStructure_L{n}_v2_ReferenceSheets.pdf'))],check=True)
d=doc();d.add_heading('Lesson 4 · prediction outline',0);d.add_paragraph('Name: __________________  Date: __________')
d.add_paragraph('Write any organelle names you expect. Star those you think animal cells share. Do this BEFORE seeing a labelled reference.');t=d.add_table(rows=1,cols=1);t.style='Table Grid';t.rows[0].height=Inches(5.3);t.rows[0].height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
d.add_page_break();d.add_heading('Lesson 4 · Organelle Job Card',0);d.add_paragraph('Complete individually during Explain. Write the structure that matches each farm job.');table(d,['Farm job','Cell structure'],[[b,''] for a,b,c,e,f in ORG]);d.add_paragraph('Keep your prediction and corrected annotation in your notebook.');d.save(keep(a.out/'CellStructure_L4_v2_BlankTasks.docx'))
d=doc();d.add_heading('Lesson 6 · my observation record',0);d.add_paragraph('Name: __________________  Group: ____  Date: ________');d.add_paragraph('Only draw what YOU see. Keep unclear features uncertain. No expected result is supplied.')
for title in ['Low power · draw 3–4 cells','Higher power · draw ONE cell']:
 d.add_heading(title,2);d.add_paragraph('Eyepiece ____ ×  Objective ____ ×  Total optical magnification ____ ×');t=d.add_table(rows=1,cols=1);t.style='Table Grid';t.rows[0].height=Inches(1.7);t.rows[0].height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
d.add_page_break();d.add_heading('Measurement and explanation',1);table(d,['Quantity','My value'],[['Measured field diameter at objective ____ ×','____ mm'],['Cell lengths across the field, along long axis','____'],['Length ≈ diameter ÷ count','____ mm'],['Length in μm = length in mm ×1000','____ μm']])
for q in ['What structures could I identify? What was unclear or not seen?','Chloroplast visibility: was my view clear enough to judge?','How does my evidence change my prediction?','My three-sentence Claim–Evidence–Reasoning explanation','My prediction for the cheek-cell practical, with a reason']:
 d.add_heading(q,2);d.add_paragraph('________________________________________________________________\n________________________________________________________________')
d.save(keep(a.out/'CellStructure_L6_v2_ObservationSheet.docx'))
d=doc();d.add_heading('Cell Structure v2 · design and review guide',0)
for ln in (ROOT/'content/cellstructure_v2/DESIGN.md').read_text().splitlines():
 if ln.startswith('## '):d.add_heading(ln[3:],1)
 elif ln and not ln.startswith('# '):d.add_paragraph(ln)
d.add_heading('Asset prompts and saved files',1);d.add_paragraph('Built-in imagegen was used. The exact four prompts, tool mode and checksums are included in image_provenance.json. The project assets live in assets/cellstructure_v2/{amina,kitchen,farm,microscope}.png. Scientific diagrams remain editable native PowerPoint objects.')
d.save(keep(a.out/'CellStructure_v2_DesignGuide.docx'))
shutil.copy2(ROOT/'assets/cellstructure_v2/provenance.json',keep(a.out/'image_provenance.json'))
# Bundle evidence and inherited detailed correction logs alongside fresh per-deck logs.
qa_files=[ROOT/'qa/CellStructure_v2_content_checks.json',ROOT/'qa/CellStructure_v2_sources.json',ROOT/'qa/CellStructure_v2_visual_review.json',ROOT/'qa/CellStructure_v2_assets_gpu.json',ROOT/'qa/CellStructure_v2_notes_font_check.json']
qa_files.extend(ROOT/'qa'/f'{name}_{suffix}.json' for name in NAMES.values() for suffix in ['validation','manifest','gpu'])
logs=list((ROOT/'build_logs').glob('*_v2.md'))
logs.extend(ROOT/'build_logs'/f'{NAMES[n].removesuffix("_v2")}.md' for n in [1,2,6])
readme=keep(a.out/'CellStructure_v2_README.txt');readme.write_text('CELL STRUCTURE · LESSONS 1, 2, 4, 6 · VERSION 2\nREVIEW DRAFTS — NOT APPROVED FOR CLASSROOM RELEASE\n\nStart with CellStructure_v2_QuickReview.pdf (17 selected slides), then review each separate PPTX and full PDF. PPTX contains editable text/science diagrams and complete teacher speaker notes. Slides 2–3 are hidden; PDFs deliberately include them.\n\nL1: 21 slides / 40 min. L2: 25 / 80. L4: 37 / 80. L6: 27 / 80.\n\nThe DesignGuide explains the repeatable method. Exact illustration prompts are in image_provenance.json. New L4 applies the same design components to content absent from the three supplied references. Source excerpts under reusable_design require the full repository (source data, theme, notes, assets, helper module and pinned requirements); the ZIP is a review package, not a standalone build environment.\n\nPrint L4 BlankTasks before Predict; labelled reference sheets only after Predict. L2 reference sheets contain the reading/table/cards. L6 sheet is blank: actual observations remain the pupil evidence.\n\nOpen checks: native PowerPoint and Presenter View; exact Calibri rendering; ARES/offline resources; teacher/source-owner decisions; L6 stain/PPE/equipment/calibration. L4 uses the source-permitted diagram-only route, not fabricated EM images. See build logs and qa evidence in the ZIP.\n')
bundle=a.out/'CellStructure_Lessons1-2-4-6_v2_Review.zip'
with zipfile.ZipFile(bundle,'w',zipfile.ZIP_DEFLATED) as z:
 for f in files:z.write(f,f.name)
 for f in qa_files:
  assert f.exists(),f;z.write(f,'qa/'+f.name)
 for f in logs:
  if f.exists():z.write(f,'build_logs/'+f.name)
 z.write(ROOT/'content/cellstructure_v2/DESIGN.md','reusable_design/DESIGN.md')
 z.write(ROOT/'scripts/lesson_visuals.py','reusable_design/lesson_visuals.py');z.write(ROOT/'scripts/build_cellstructure_v2.py','reusable_design/build_cellstructure_v2.py')
for f in files+[bundle]:
 if f.suffix in ['.zip','.docx','.pptx']:
  with zipfile.ZipFile(f) as z:assert z.testzip() is None
print(json.dumps({'bundle':str(bundle),'files':len(files),'bytes':bundle.stat().st_size,'sha256':hashlib.sha256(bundle.read_bytes()).hexdigest()},indent=2))
