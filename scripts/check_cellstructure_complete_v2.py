#!/usr/bin/env python3
"""Check the complete selected unit: notes/font/timing, source phases, text retained in PDF and protected reveals."""
import json,re,subprocess,hashlib,zipfile
from pathlib import Path
from pptx import Presentation
from build_cellstructure_remaining_v2 import ALL_NAMES,REMAINING_NAMES,ROOT
reports=[]
def render_folder(n):return ROOT/'qa/renders'/('cellstructure_complete_v2' if n in REMAINING_NAMES else 'cellstructure_v2')
def words(t):return set(re.findall(r'[a-z]{4,}',t.lower()))
for n,name in sorted(ALL_NAMES.items()):
 p=ROOT/'outputs/drafts/v2'/f'{name}.pptx';pdf=render_folder(n)/f'{name}.pdf';d=Presentation(p);errors=[];times=[];titles=[]
 pages=subprocess.check_output(['pdftotext','-layout',str(pdf),'-'],text=True).split('\f')
 for i,s in enumerate(d.slides):
  tf=s.notes_slide.notes_text_frame;txt=tf.text
  for para in tf.paragraphs:
   if para.font.name!='Calibri' or any(r.font.name!='Calibri' for r in para.runs):errors.append(f'{i+1}: notes font')
  t=re.findall(r'^TIME: ([\d.]+) minutes$',txt,re.M)
  if len(t)!=1:errors.append(f'{i+1}: TIME')
  else:times.append(float(t[0]))
  texts=[]
  for sh in s.shapes:
   if sh.has_text_frame:texts.append(sh.text)
   if sh.has_table:texts.extend(c.text for row in sh.table.rows for c in row.cells)
  missing=words('\n'.join(texts))-words(pages[i])
  if missing:errors.append({'slide':i+1,'missing_pdf_words':sorted(missing)})
  titles.append('\n'.join(texts))
 total=sum(times);expected=40 if n in [1,3,11] else 80
 if abs(total-expected)>1e-6:errors.append(f'Total {total} != {expected}')
 hidden=[i+1 for i,s in enumerate(d.slides) if s._element.get('show')=='0']
 if hidden!=[2,3]:errors.append('Hidden slides')
 for phase in ['PREDICT','OBSERVE','EXPLAIN','DRIVING QUESTION BOARD','MODEL BUILDING']:
  if not any(phase in text for text in titles):errors.append('Missing phase '+phase)
 if n==7:
  own=next(i for i,t in enumerate(titles) if 'Your turn · observe, draw, label' in t);compare=next(i for i,t in enumerate(titles) if 'Compare AFTER recording' in t)
  if compare<=own:errors.append('Expected comparison precedes pupil recording')
 if n==8:
  if not any('It does not heal a finger cut.' in t for t in titles):errors.append('Sperm correction absent')
 if n==12:
  assessment=next(i for i,t in enumerate(titles) if 'Final Explanation · independent work' in t)
  if times[assessment]!=25:errors.append('Assessment start is not 25 minutes')
 with zipfile.ZipFile(p) as z:
  if z.testzip():errors.append('PPTX CRC')
 report={'lesson':n,'name':name,'slides':len(d.slides),'minutes':round(total,6),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'status':'PASS' if not errors else 'FAIL','errors':errors}
 reports.append(report)
result={'lessons':reports,'total_slides':sum(r['slides'] for r in reports),'new_slides':sum(r['slides'] for r in reports if r['lesson'] in REMAINING_NAMES),'status':'PASS' if not any(r['errors'] for r in reports) else 'FAIL','limits':'Structural/source-sequence/PDF-word retention checks. This does not replace rendered inspection, teacher science review, actual microscope trials or native PowerPoint/Presenter View.'}
(ROOT/'qa/CellStructure_complete_v2_checks.json').write_text(json.dumps(result,indent=2));print(json.dumps(result,indent=2));raise SystemExit(result['status']!='PASS')
