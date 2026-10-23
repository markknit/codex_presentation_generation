#!/usr/bin/env python3
"""Build the Cell Structure L1 review pilot from the approved visual example and V2 sources."""
from copy import deepcopy
from pathlib import Path
import hashlib
import json
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT=Path(__file__).resolve().parents[1]
UNIT=ROOT/'inputs/v2/Biology/SS1.1_Cell_Structure'
TEMPLATE=ROOT/'references/templates/Bio_G10_CellStructure_L1_The_Healing_Cut.pptx'
NAME='Bio_G10_CellStructure_L1_The_Healing_Cut'

def set_text(shape,text,size=None):
    tf=shape.text_frame
    p=tf.paragraphs[0]
    paragraph_style=deepcopy(p._p.pPr) if p._p.pPr is not None else None
    original=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    tf.clear()
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if paragraph_style is not None:
            old=p._p.pPr
            if old is not None:p._p.remove(old)
            p._p.insert(0,deepcopy(paragraph_style))
        r=p.add_run();r.text=line
        if original is not None:r._r.insert(0,deepcopy(original))
        r.font.name='Calibri'
        if size:r.font.size=Pt(size)
    tf.word_wrap=True

def add_text(slide,text,x,y,w,h,size=13,color='555555',bold=False):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    shape.text_frame.margin_left=shape.text_frame.margin_right=0
    shape.text_frame.margin_top=shape.text_frame.margin_bottom=0
    set_text(shape,text,size)
    for p in shape.text_frame.paragraphs:
        for r in p.runs:r.font.color.rgb=RGBColor.from_string(color);r.font.bold=bold
    return shape


def main():
    unit=json.loads((UNIT/'Biology_Cell_Structure_data.json').read_text());lesson=unit['LESSONS'][0]
    deck=Presentation(TEMPLATE)
    slides=list(deck.slides)
    # Complete phase names on student and hidden teacher slides.
    for slide in slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:continue
            if sh.text=='QUESTION BOARD':
                set_text(sh,'DRIVING QUESTION BOARD',10 if sh.width<Inches(3) else 12)
                if sh.width<Inches(3.8):
                    sh.width=Inches(3.8)
                    # Matching coloured background label, when present.
                    for bg in slide.shapes:
                        if bg!=sh and abs(bg.left-sh.left)<10 and abs(bg.top-sh.top)<10 and bg.height==sh.height:bg.width=sh.width
            elif sh.text=='Question Board':set_text(sh,'Driving Question Board',10 if sh.width<Inches(2.5) else 13)
    # Wide title-phase strip avoids a wrapped label.
    set_text(slides[0].shapes[11],'Driving Question\nBoard',18)
    slides[0].shapes[11].height=Inches(.72)
    slides[0].shapes[12].top=slides[0].shapes[13].top=Inches(4.78)
    add_text(slides[0],'REVIEW DRAFT · 23 OCTOBER 2026',.8,6.2,6,.3,11,'FCBA63')
    # Same routine text will be reused in subsequent approved builds.
    set_text(slides[1].shapes[12],'Help pupils make sense of evidence. In Lesson 1, keep causal ideas tentative; later lessons test them.',13)
    # Keep source reveal order; the formal question is first displayed on slide 19.
    set_text(slides[3].shapes[44],'Our driving question\nWe will reveal the full question after you post what you notice and wonder.',17)
    set_text(slides[3].shapes[42],'So far\nThis is our first lesson. We begin with Amina’s cut finger and our own ideas.',17)
    # Three concise goals, then plain-language process vocabulary (not cellular answers).
    for j in [3,4,5]:slides[4].shapes[j].top-=Inches(.2)
    for j in [6,7,8]:slides[4].shapes[j].top-=Inches(.45)
    for j in [9,10,11]:slides[4].shapes[j].top-=Inches(.7)
    add_text(slides[4],'Key words: observe = describe what you notice · predict = say what you think may happen',.65,5.65,12,.4,16,'333333')
    add_text(slides[4],'model = a drawing that shows your current ideas',.65,6.12,12,.35,16,'333333')
    # Glance timings include the opening; no extra unbudgeted introduction.
    set_text(slides[2].shapes[9],'Opening, Amina’s story, then 2 full minutes of independent writing.',11.5)
    set_text(slides[2].shapes[27],'Summary',12)
    set_text(slides[2].shapes[35],'Wound Healing Process — preview first',11)
    set_text(slides[2].shapes[38],'Human body organisation — not today',11)
    set_text(slides[2].shapes[41],'Cell parts & functions — not today',11)
    # Resource links are copied exactly from V2; never claimed verified on ARES.
    links=[lesson['resourceLinks']['predict']['reading']['direct_url'],lesson['resourceLinks']['predict']['video']['direct_url'],lesson['resourceLinks']['explain']['video']['direct_url']]
    for ix,url in zip([35,38,41],links):
        for p in slides[2].shapes[ix].text_frame.paragraphs:
            for r in p.runs:
                r.hyperlink.address=None
                r.font.color.rgb=RGBColor(255,255,255)
                r.font.underline=False
    slides[2].shapes[35].click_action.hyperlink.address=links[0]
    slides[2].shapes[38].click_action.hyperlink.address=links[1]
    slides[2].shapes[41].click_action.hyperlink.address=links[2]
    # Source-supported baseline tasks get visible formative labels.
    for n in [7,15,16,18,19]:
        add_text(slides[n-1],'CHECK FOR UNDERSTANDING',8.6,.45,4.1,.35,10,'333333' if n!=19 else 'FFFFFF',True)
    # Protect later panels; explicitly identify the inherited artwork as illustrations.
    captions=[('1. The cut\njust now'),('2. A clot forms\na little later'),('3. A scab\ntwo days later'),('4. New skin\nover several days')]
    for n in range(9,13):
        s=slides[n-1]
        for sh in s.shapes:
            if sh.has_text_frame and re.match(r'^[1-4]\. ',sh.text):
                panel=int(sh.text[0]);revealed=n-8
                set_text(sh,captions[panel-1] if panel<=revealed else f'{panel}. What changes next?',16)
        add_text(s,'Illustrative drawings — not photographs. This story is not a fixed healing timetable.',.65,6.25,12,.35,12,'555555')
    # Keep the notice faithful to the shown story, without an unsupported healing deadline.
    set_text(slides[17].shapes[9],'I notice…\n\n…the bleeding slows and the skin changes over time.',21)
    set_text(slides[17].shapes[12],'Your turn\nYellow: one thing you NOTICE.\nBlue: one real question you WONDER.\nPost on our Driving Question Board.',20)
    set_text(slides[20].shapes[19],'Next lesson: What is a cell? Keep today’s predictions — we will return to them.',18)
    # Notes use exactly one numeric TIME per slide. Hidden preparation slides have zero class time.
    times=[.25,0,0,.5,.5,1.25,2.5,2,1,1,1,2,3,1.5,1.5,2,2,7,3,5,3]
    assert sum(times)==40
    notes=[s.notes_slide.notes_text_frame.text for s in slides]
    for i,text in enumerate(notes):
        # Existing source-derived notes are retained; headings are normalised below.
        text=re.sub(r'^TIME:.*\n?', '',text,flags=re.M)
        text=text.replace('LOOK FOR (formative check)','LOOK FOR').replace('TEACHER REFLECTION (after the lesson)','TEACHER REFLECTION')
        text=text.replace('5 minutes for all four pictures (slides 9–12)','5 minutes across slides 9–12')
        notes[i]=text.strip()
    # Preserve source narration in the fixed label order.
    notes[0]='''PURPOSE OF THIS LESSON
Launch the phenomenon and record starting ideas, not a cellular explanation. Read the entire 12-lesson source set before teaching. Cell theory belongs to Lesson 2; organelle functions to Lessons 4–5; specialised wound-healing roles are developed later.

BEFORE CLASS
Prepare a poster with What do we notice? and What do we wonder? Cover the formal driving question until slide 19. Provide one yellow and one blue note per pupil, notebooks, writing materials and a visible timer. If sticky notes are unavailable, use paper slips and tape. Have the Summary Table ready. The four finger drawings are illustrative; no injury demonstration or practical is used.

DO
Use Presenter View. Keep slides 2–3 hidden. Spend 15 seconds introducing the title; the full opening is included in the 5-minute Predict budget. Do not show or preview cell anatomy resources today.

SAY
Today we begin an investigation. Your own ideas and questions are our starting point.

IF…
If projection is unavailable, narrate the story and read the picture descriptions from slides 9–12. No online resource is required.

SCIENCE NOTE FOR YOU
The story illustrates a sequence, not a guaranteed healing schedule. Do not infer microscopic causes from these drawings. This is a review draft; native PowerPoint, teacher approval and Calibri fidelity remain unchecked.'''
    notes[1]+='''

ANSWER KEY
A prediction is an initial idea; an observation describes what is seen or described in the evidence; a model represents current thinking. These process definitions do not reveal the unit's biological explanation.'''
    notes[2]='''DO
Prepare all materials before pupils arrive. Timing: Predict including opening 5 min; Observe 10 min; Explain 7 min; Driving Question Board 10 min; Model Building 5 min; Summary Table 3 min. Total 40 min. Slides 2 and 3 have zero teaching time.

SAY
No new biological content is to be delivered today. Do not show cell labels, cell theory or wound-healing mechanisms.

IF…
No projector: narrate each stage and ask pupils to record the changes they hear. No sticky notes: use paper and tape. Do not replace the lesson with an unpreviewed video.

SCIENCE NOTE FOR YOU
Resource audit (names and URLs copied from V2; local availability and actual content have not been verified):
• Wound Healing Process (PLIX), CK-12 — Preview first, not an exact teaching match. Teacher preparation only today: it may disclose mechanisms before pupils investigate.
• Organization in the human body, Khan Academy — Do not use in Lesson 1; it moves to organisation before the launch inquiry.
• Cell parts and their functions; Cell structures and their functions; Comparing animal and plant cells; Plant vs animal cells; Overview of animal and plant cells; Apply: cell parts and functions — Do not use today; these reveal later content. Reconsider during preparation of the relevant later lesson, after previewing.
The three buttons retain the source ARES URLs. No link has been invented or presented as tested.'''
    # Targeted source-derived interventions and task criteria.
    extras={
     4:('IF…','If pupils read ahead, say: Keep that question. We will reveal our shared question after your own notice and wonder notes.'),
     5:('SAY','Observe means describe what you notice. Predict means give your starting idea. A model is a drawing of your current thinking.'),
     6:('IF…','If the injury story is uncomfortable, pupils may keep their eyes open or look away and listen. Use the non-graphic drawings; do not use injury photographs.'),
     7:('IF…','If a pupil cannot begin, preserve the silent writing period. Note the blank response for later support; do not supply a cell name or mechanism.'),
     8:('IF…','If only one partner speaks, ask the quieter partner to read their own written prediction. Keep the full 90-second sharing window.'),
     9:('LOOK FOR','Pupils describe the red line and blood from the narrated drawing, rather than naming an unseen mechanism. Read at least four notebook additions across the sequence.'),
     10:('LOOK FOR','Pupils distinguish the visible slowing of bleeding/darker material from an explanation of why it happens.'),
     11:('LOOK FOR','Pupils compare the new visible scab with the earlier appearance. They should not copy a future panel caption.'),
     12:('IF…','If a pupil names a cell type, record it as their idea without confirming it. Say: We will test that idea in later lessons.'),
     13:('IF…','If pupils look to you for a verdict, say: I am recording your words. We will investigate which explanations the evidence supports.'),
     14:('LOOK FOR','Listen for causal suggestions, keeping these tentative. No pupil is required to provide the later cellular explanation.'),
     15:('IF…','If pupils offer only one cause, use the source follow-up: Why do you think that? Do not teach the separate cell roles.'),
     16:('WHY THIS WORKS','The question moves from cause to the possibility of several structures doing different jobs; it does not establish their identities.'),
     17:('SCIENCE NOTE FOR YOU','This is the source bridging statement. Do not name platelets, neutrophils, macrophages or fibroblasts unless recording a pupil’s own prior idea; do not confirm their roles today.'),
     18:('LOOK FOR','Check that each pupil produces one observation and one genuine question. Use the actual healing sequence as the observation source, not a textbook mechanism.'),
     19:('IF…','If a note is an explanation presented as an observation, ask: Which change did you actually see or hear described? Keep the proposed cause as an idea to investigate.'),
     20:('IF…','If a pupil asks for a cell diagram, say: Draw whatever comes to mind. This records your starting idea, not a correct-answer test.'),
     21:('IF…','If the explanation column is blank, prompt: What can we explain yet, and what are we still wondering? The source explicitly accepts that the cellular explanation is still unknown.')}
    for n,(label,text) in extras.items():notes[n-1]+='\n\n'+label+'\n'+text
    # Add source-grounded answer-key criteria for the actual written tasks, not a supplied baseline diagram.
    notes[6]+='\n\nANSWER KEY\nNo single correct biological response is expected. Preserve each pupil’s own initial idea, uncorrected, for comparison in Lesson 12.'
    notes[7]+='\n\nANSWER KEY\nPairs report a similarity and a difference between their own predictions; no mechanism is confirmed.'
    for n in [9,10,11,12]:
        notes[n-1]+='\n\nANSWER KEY\nThe source sequence is: blood appears; bleeding slows and a clot forms; a scab develops; new skin appears over several days. Accept only the stage already revealed. These are illustrative observations, not a microscopic explanation.'
    notes[17]+='\n\nIF…\nIf pupils write what they think you want, repeat the source prompt: Your wonder note should be a real question you actually have.\n\nANSWER KEY\nA notice refers to an appearance or change in the sequence. A wonder is the pupil’s own open question; there is no single correct question.'
    notes[18]+='\n\nANSWER KEY\nThere is no required first reaction. Record the actual questions; do not supply a unit answer.'
    notes[19]+='\n\nANSWER KEY\nThere is no correct baseline cell drawing. Accept the pupil’s own sketch and any labels they already know. Do not show an exemplar. Compare with later revisions.'
    # Fixed note order; repeated labels are merged without discarding content.
    order=['PURPOSE OF THIS LESSON','BEFORE CLASS','DO','SAY','WAIT TIME','LOOK FOR','IF…','WHY THIS WORKS','ANSWER KEY','SCIENCE NOTE FOR YOU','TEACHER REFLECTION']
    for i,(s,text) in enumerate(zip(slides,notes)):
        sections={};heading='DO';phase=''
        for line in text.splitlines():
            key=line.strip()
            if key in order:heading=key;continue
            if re.match(r'^(PREDICT PHASE|OBSERVE PHASE|EXPLAIN PHASE|DRIVING QUESTION BOARD \(|MODEL BUILDING PHASE)',key):phase=key;continue
            if key:sections.setdefault(heading,[]).append(line)
        if i==5:phase='PREDICT PHASE (5 minutes total, including title, map and goals)'
        if i==19:phase='MODEL BUILDING PHASE (5 minutes total)'
        body=(phase+'\n' if phase else '')+f'TIME: {times[i]:g} minutes\n'
        body+='\n'.join('\n'+key+'\n'+'\n'.join(sections[key]) for key in order if key in sections)
        s.notes_slide.notes_text_frame.text=body
    # Keep all explicitly authored text Calibri, including inherited runs.
    for s in slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:r.font.name='Calibri'
        for p in s.notes_slide.notes_text_frame.paragraphs:
            for r in p.runs:r.font.name='Calibri'
    deck.core_properties.title='The Healing Cut — Biology Grade 10, Lesson 1'
    deck.core_properties.subject='Cell Structure • review pilot'
    deck.core_properties.comments='Review draft. Not teacher-approved. Lessons 2–3 await pilot gate.'
    target=ROOT/'outputs/drafts'/f'{NAME}.pptx';target.parent.mkdir(parents=True,exist_ok=True);deck.save(target)
    files=sorted(UNIT.glob('*'))+[TEMPLATE, ROOT/'references/instructions/SeaVuria_Science_Lesson_PPT_Build_Manual.docx']
    manifest={'unit':'Biology / SS1.1_Cell_Structure','requested_lessons':[1,2,3],'built_lessons':[1]+[n for n in (2,3) if list((ROOT/'outputs/drafts').glob(f'Bio_G10_CellStructure_L{n}_*.pptx'))],'deferred_lessons':[n for n in (2,3) if not list((ROOT/'outputs/drafts').glob(f'Bio_G10_CellStructure_L{n}_*.pptx'))],'defer_reason':'See STATUS.md for draft continuation authorisation and remaining review gates','lesson_titles':[{'number':x['number'],'title':x['title'],'duration':x['duration']} for x in unit['LESSONS']],'source_label_conflict':'Folder/META/lesson labels say SS1.1; UNIT metadata says Sub-Strand 1.3. Deck uses Cell Structure without a numerical sub-strand claim; owner review required.','inputs':[{'path':str(p.relative_to(ROOT)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in files if p.is_file()],'output_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'slide_times_minutes':times,'total_minutes':sum(times)}
    (ROOT/'qa/Bio_G10_CellStructure_manifest.json').write_text(json.dumps(manifest,indent=2))
    print(target)

if __name__=='__main__':main()
