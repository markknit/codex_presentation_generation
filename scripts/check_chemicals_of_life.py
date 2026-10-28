#!/usr/bin/env python3
"""Narrow source-order checks; visual, scientific and native checks remain separate."""
import json, hashlib
from pptx import Presentation
from build_chemicals_of_life import ROOT,NAMES,OUT

def text(s):
    values=[]
    for sh in s.shapes:
        if sh.has_text_frame:values.append(sh.text)
        if sh.has_table:values.extend(c.text for row in sh.table.rows for c in row.cells)
    return '\n'.join(values)

errors=[];decks=[]
for n,name in NAMES.items():
    p=OUT/(name+'.pptx');prs=Presentation(p);texts=[text(s) for s in prs.slides];notes=[s.notes_slide.notes_text_frame.text for s in prs.slides]
    val=json.loads((ROOT/'qa'/f'{name}_validation.json').read_text())
    gpu=json.loads((ROOT/'qa'/f'{name}_gpu.json').read_text())
    digest=hashlib.sha256(p.read_bytes()).hexdigest()
    if val['structural_status']!='PASS' or val['sha256']!=digest:errors.append(f'L{n}: structural result or hash')
    if gpu['deck_sha256']!=digest or len(gpu['slides'])!=len(prs.slides):errors.append(f'L{n}: CUDA report or hash')
    for label in ['PREDICT','OBSERVE','EXPLAIN','DRIVING QUESTION BOARD','MODEL BUILDING']:
        if not any(label in x for x in texts):errors.append(f'L{n}: missing phase {label}')
    if n in [2,3,4,5] and not any('SAFETY' in x for x in texts):errors.append(f'L{n}: missing safety slide')
    if n==1:
        predict=next(i for i,x in enumerate(texts) if 'Predict the shape' in x)
        reveal=next(i for i,x in enumerate(texts) if 'Read Wanjiru’s case data' in x)
        if predict>=reveal:errors.append('L1: graph reveal precedes prediction')
        if any('8.0' in x for x in texts[:reveal] if 'UNIT MAP' not in x):errors.append('L1: peak value displayed before data')
    if n==3:
        observation=next(i for i,x in enumerate(texts) if 'Benedict’s test · record the colour' in x)
        guide=next(i for i,x in enumerate(texts) if 'Compare with the test guide' in x)
        if observation>=guide:errors.append('L3: test guide before recording')
    if n==5 and not any('pooled' in x.lower() and 'not endorse' in x.lower() for x in notes):errors.append('L5: pooled-saliva caution missing')
    if n==6 and not any('five-part' in x.lower() and 'supervised' in x.lower() for x in notes):errors.append('L6: assessment follow-up missing')
    decks.append({'lesson':n,'slides':len(prs.slides),'sha256':digest,'minutes':sum(val['slide_minutes']),'cuda_previews':len(gpu['slides'])})
out={'status':'PASS' if not errors else 'FAIL','errors':errors,'decks':decks,'limits':'Only checks selected sequencing, timing reports, hashes, phase labels and safety flags. It does not establish complete source coverage, scientific correctness, visual quality, human safety, native PowerPoint behaviour or teacher approval.'}
(ROOT/'qa/ChemicalsOfLife_checks.json').write_text(json.dumps(out,indent=2))
print(json.dumps(out,indent=2));raise SystemExit(bool(errors))
