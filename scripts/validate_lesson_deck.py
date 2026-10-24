#!/usr/bin/env python3
"""Check review-deck package, notes, timing, text geometry and required teaching metadata."""
import argparse,hashlib,json,re,subprocess,zipfile
from pathlib import Path
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
p=argparse.ArgumentParser();p.add_argument('deck',type=Path);p.add_argument('--minutes',type=float,required=True);a=p.parse_args()
d=Presentation(a.deck);errors=[];times=[];hidden=[];checks=[]
order=['PURPOSE OF THIS LESSON','BEFORE CLASS','DO','SAY','WAIT TIME','LOOK FOR','IF…','WHY THIS WORKS','ANSWER KEY','SCIENCE NOTE FOR YOU','TEACHER REFLECTION']
for i,s in enumerate(d.slides,1):
    n=s.notes_slide.notes_text_frame.text if s.notes_slide.notes_text_frame else ''
    if not n:errors.append(f'slide {i}: missing notes')
    tt=re.findall(r'^TIME: ([\d.]+) minutes$',n,re.M)
    if len(tt)!=1:errors.append(f'slide {i}: expected one numeric TIME')
    else:times.append(float(tt[0]))
    if s._element.get('show')=='0':hidden.append(i)
    pos=[n.find('\n'+k+'\n') for k in order if '\n'+k+'\n' in n]
    if pos!=sorted(pos):errors.append(f'slide {i}: note order')
    texts=[]
    for sh in s.shapes:
        if sh.has_text_frame:texts.append(sh.text)
        if sh.width>0 and sh.height>0 and (sh.left<-20000 or sh.top<-20000 or sh.left+sh.width>d.slide_width+20000 or sh.top+sh.height>d.slide_height+20000):errors.append(f'slide {i}: out of bounds {sh.name}')
        frames=[sh.text_frame] if sh.has_text_frame else []
        if sh.has_table:frames.extend(c.text_frame for row in sh.table.rows for c in row.cells)
        for tf in frames:
            for para in tf.paragraphs:
                for r in para.runs:
                    if r.font.name!='Calibri':errors.append(f'slide {i}: non-Calibri text')
    t='\n'.join(texts)
    if '© 2026 SeaVuria & ARES Education' not in t:errors.append(f'slide {i}: missing footer')
    if 'CHECK FOR UNDERSTANDING' in t:
        checks.append(i)
        for k in ['LOOK FOR','IF…']:
            if '\n'+k+'\n' not in n:errors.append(f'slide {i}: missing {k}')
    if i not in [2,3] and re.search(r'\bDQB\b',t):errors.append(f'slide {i}: abbreviated Driving Question Board')
if hidden!=[2,3]:errors.append('only slides 2 and 3 should be hidden')
if abs(sum(times)-a.minutes)>1e-6:errors.append(f'total {sum(times)} != {a.minutes}')
if abs(d.slide_width/d.slide_height-16/9)>.01:errors.append('not 16:9')
with zipfile.ZipFile(a.deck) as z:
    if z.testzip():errors.append('package CRC failure')
report={'deck':str(a.deck),'sha256':hashlib.sha256(a.deck.read_bytes()).hexdigest(),'structural_status':'PASS' if not errors else 'FAIL','errors':errors,'slides':len(d.slides),'hidden_slides':hidden,'slide_minutes':times,'total_minutes':sum(times),'formative_check_slides':checks,'render_font':subprocess.check_output(['fc-match','Calibri'],text=True).strip(),'release_status':'DRAFT — NOT APPROVED','pending':['Native PowerPoint and Presenter View','Teacher/source-owner approval','Calibri fidelity','ARES link/offline/content verification'],'limits':'Geometry detects canvas bounds, not all text overflow. Rendered inspection and science review are separate.'}
(ROOT/'qa'/f'{a.deck.stem}_validation.json').write_text(json.dumps(report,indent=2));print(json.dumps(report,indent=2));raise SystemExit(bool(errors))
