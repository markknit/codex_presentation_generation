#!/usr/bin/env python3
"""Create editable classroom records and a local review package, not a release."""
import hashlib, json, shutil, subprocess, zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.table import WD_ROW_HEIGHT_RULE
from build_chemicals_of_life import ROOT, SOURCE, DATA, NAMES, OUT

TARGET=ROOT/'outputs/drafts/chemicals_of_life_review'
TARGET.mkdir(parents=True,exist_ok=True)

def doc(title, sub):
    d=Document(); sec=d.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
    style=d.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(11)
    d.add_heading(title,0);d.add_paragraph(sub);return d
def grid(d,head,rows):
    t=d.add_table(rows=1, cols=len(head));t.style='Table Grid';t.autofit=False
    for c,v in zip(t.rows[0].cells,head):c.text=str(v)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for row in t.rows:
        row.height=Inches(.28);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
    return t

d=doc('Chemicals of Life · case and evidence cards','Grade 10 Biology · review draft · Source case, not a diagnostic tool')
d.add_heading('Wanjiru’s case',1);d.add_paragraph('After fasting, Wanjiru ate ugali (300 g), sukuma wiki (100 g) and beans (100 g) between 0 and 15 minutes. The case source provides the readings below; their independent provenance has not been verified.')
grid(d,['Time after meal start (min)','Blood glucose concentration (mmol/L)'],list(zip([0,15,30,60,90,120,150,180],[4.5,5.5,7.0,8.0,6.5,5.2,4.7,4.5])))
d.add_paragraph('The source gives a fasting reference range of 4.0–5.5 mmol/L. This case is not a diagnosis.')
d.add_heading('Source meal estimates',1);grid(d,['Food','Approximate source information'],[['Ugali','70 g carbohydrate'],['Beans','9 g protein; 23 g carbohydrate'],['Sukuma wiki','3 g protein; fibre, vitamins A and C; iron and calcium'],['Cooking oil','10 g fat']]);d.add_paragraph('These estimates are given in the lesson source; they are not class measurements and depend on food preparation.')
d.add_page_break();d.add_heading('Element cards · cut or copy',1)
grid(d,['Symbol','Source example'],[['C, H, O','Starch and glucose'],['N, S','Protein; some proteins contain sulfur'],['Ca','Milk'],['Fe','Sukuma wiki or nyama choma'],['P, Mg, Na, K','Teacher: supply a checked source card or example before this station.']])
d.add_paragraph('The source names all eleven symbols but does not supply a complete set of individual evidence cards. The teacher must supply the missing card details; pupils should cite the information they were actually given.')
d.save(TARGET/'ChemicalsOfLife_Case_and_Cards.docx')

d=doc('Chemicals of Life · investigation record','Name: ____________________   Class: __________   Date: __________\nWrite actual observations only. If an investigation was not performed, state that.')
d.add_heading('Lesson 3 · iodine and Benedict’s',1)
grid(d,['Sample','Test','Actual colour / appearance','What might it indicate?'],[[x,'Iodine','',''] for x in ['Ugali','Maize flour','Sukuma wiki','Milk','Nyama choma']]+[[x,'Benedict’s','',''] for x in ['Banana','Ugali','Milk','Glucose solution']])
d.add_page_break();d.add_heading('Lesson 4 · protein and lipid tests',1)
grid(d,['Sample','Test','Actual appearance','Interpretation'],[[x,'Biuret','',''] for x in ['Beans','Milk','Nyama choma','Ugali']]+[[x,'Grease spot / emulsion','',''] for x in ['Groundnuts','Mandazi','Milk','Ugali']])
d.add_page_break();d.add_heading('Lesson 5 · starch and control',1)
d.add_paragraph('Do not perform until the teacher and school have approved the specimen and practical safety plan. Control chosen: __________________________')
grid(d,['Time (min)','Iodine: tube A actual colour','Iodine: tube B actual colour'],[[0,'',''],[5,'',''],[10,'','']])
grid(d,['At 10 min','Benedict’s: actual colour'],[['Tube A',''],['Tube B','']])
d.add_paragraph('What can the control help us compare? __________________________________________\nWhat are the limitations? _____________________________________________________')
d.save(TARGET/'ChemicalsOfLife_Investigation_Records.docx')

d=doc('Lesson 6 · role cards','Cut or copy. This is a conceptual model, not a measurement of Wanjiru’s body.')
grid(d,['Role','What to represent'],[['Glucose','Use counters to represent an amount of glucose.'],['Small intestine wall','Show glucose moving into the blood.'],['Blood','Use a rope loop as a simplified part of the circulation.'],['Pancreas','Use insulin cards to represent a hormone signal.'],['Body cells','Show glucose use; include oxygen in the model.'],['Liver and muscle','Represent storage as glycogen.'],['Observer','Record the sequence; compare its broad pattern with the case graph.']])
d.add_paragraph('A card is not a measured molecule count. The role-play is not evidence of Wanjiru’s actual hormone concentrations or organ activity.')
d.save(TARGET/'ChemicalsOfLife_L6_Role_Cards.docx')

d=doc('Final Explanation · student assessment','Review draft · arrange supervised follow-up time after Lesson 6. Name: __________________  Date: __________')
d.add_paragraph('Use the case table and your own lesson evidence. If a practical was not conducted, say so; do not invent results. Explain which claims are supported and where evidence is limited.')
grid(d,['Time (min)','Glucose (mmol/L)'],list(zip([0,15,30,60,90,120,150,180],[4.5,5.5,7.0,8.0,6.5,5.2,4.7,4.5])))
sections=DATA['FINAL_EXPLANATION']['sections']
for i,section in enumerate(sections):
    if i:d.add_page_break()
    d.add_heading(section['title'],1)
    prompt=section['prompt']
    if i==0:prompt=prompt.replace('Give the approximate amount of each nutrient in the whole meal.','Give the approximate amounts supplied on the meal card; identify any amounts that are not supplied.')
    if i==1:prompt=prompt.replace('Ugali gave no change with Benedict\'s solution before digestion. What does this result, together with what you know about ugali and the glucose in Wanjiru\'s blood, tell you about what must happen to the starch before glucose can enter the blood?','If your ugali sample gave no change with Benedict’s before digestion, what might that suggest? Explain what must happen to starch before glucose can enter the blood, and state the limits of this test.')
    if i==2:prompt=prompt.replace('Give the results at 0, 5 and 10 minutes and for the control.','Give your recorded results at 0, 5 and 10 minutes and for the control, or state that no investigation was performed.')
    if i==3:prompt=prompt.replace('Explain why you think carbohydrates, and not proteins or lipids, are the main cause of the spike. Use evidence from the meal and from the tests.','Explain why carbohydrate is a plausible contributor to the rise. Use the meal card and your test evidence, and explain what those sources cannot establish.')
    if i==4:prompt=prompt.replace('State when the level was first back in the normal fasting range.','State when the first listed reading after the peak was back in the supplied fasting range.')
    d.add_paragraph(prompt);d.add_paragraph('Your answer: ____________________________________________________________________\n\n________________________________________________________________________________\n\n________________________________________________________________________________\n\n________________________________________________________________________________\n\n________________________________________________________________________________')
d.add_paragraph('Source rubric criteria: chemicals of life; use of evidence; enzymes and digestion; explaining the rise and fall; scientific reasoning. Teacher review must adapt the supplied rubric to the actual practical route and corrected science.')
d.save(TARGET/'ChemicalsOfLife_Final_Explanation_Student.docx')

d=doc('Teacher review notes · Chemicals of Life','Review draft — do not release for classroom use before native PowerPoint, teacher and practical safety approval.')
d.styles['Normal'].font.size=Pt(10)
for heading,text in [
 ('Case and inference','The supplied glucose values and nutrient estimates are case inputs, not independently verified records. The graph does not measure insulin or identify the causal contribution of individual foods. Do not diagnose or request personal medical disclosures.'),
 ('Food tests','Use actual observations, including unclear results. Iodine is a starch indicator; heated Benedict’s detects reducing sugars, not glucose uniquely. Sample preparation, concentrations, sensitivity and local safety plans affect interpretation.'),
 ('Lesson 5 practical','Human saliva is a specimen requiring school approval and a specific local risk assessment. The source suggests pooling; this draft does not endorse it. Choose and document an approved control and any approved alternative with the source owner. Do not fabricate data if the activity is not performed.'),
 ('Science corrections','Salivary amylase begins starch breakdown to smaller sugars including maltose; other enzymes including maltase yield glucose. Many enzymes are proteins, but not all; amylase is a protein. Temperature optima differ; cold may slow activity without eliminating it. Insulin-responsive tissue uptake and liver regulation are not identical.'),
 ('Assessment','The supplied student paper and teacher key assume particular test results and contain overgeneralised temperature and enzyme claims. The adapted student assessment asks for actual results; review and adapt the original rubric and teacher key before use. The five-part paper requires supervised time beyond the 80-minute Lesson 6 plan.'),
 ('Coverage gap','The unit learning outcomes mention investigations for vitamin C, enzymes in living tissues, and factors affecting enzyme activity. The supplied six-lesson sequence does not clearly provide a separate vitamin C or catalase investigation; obtain source-owner direction rather than inventing lessons or claiming those outcomes were demonstrated.'),
 ('Verification','All files are review drafts. Check every note and source decision, reagent SDS/concentrations, flame and ethanol controls, allergy plan, waste, human-specimen handling, timings, ARES resources, fonts, native PowerPoint and Presenter View. Obtain teacher and source-owner sign-off before release.')]:
    d.add_heading(heading,1);d.add_paragraph(text)
d.save(TARGET/'ChemicalsOfLife_Teacher_Review_Notes.docx')

# Writer export and PDF rasterisation are CPU-only; all preview resampling and
# pixel metrics are on CUDA. The report is bound to the actual DOCX bytes.
printdir=ROOT/'qa/renders/chemicals_of_life/print_final'
printdir.mkdir(parents=True,exist_ok=True)
docs=sorted(TARGET.glob('*.docx'))
for p in docs:
    subprocess.run(['soffice','-env:UserInstallation=file://'+str(ROOT/'.cache/lo-chemicals-print'),'--headless','--convert-to','pdf','--outdir',str(printdir),str(p)],check=True,stdout=subprocess.DEVNULL)
    for old in printdir.glob(p.stem+'-*.png'):old.unlink()
    subprocess.run(['pdftoppm','-png','-r','100',str(printdir/(p.stem+'.pdf')),str(printdir/p.stem)],check=True)
    shutil.copy2(printdir/(p.stem+'.pdf'),TARGET/(p.stem+'.pdf'))
images=sorted(printdir.glob('*.png'))
subprocess.run([str(ROOT/'.venv/bin/python'),str(ROOT/'scripts/gpu_visual_qa.py'),*map(str,images),'--out',str(printdir/'gpu')],check=True)
report=json.loads((printdir/'gpu/gpu_report.json').read_text())
report['documents']=[{'name':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in docs]
(ROOT/'qa/ChemicalsOfLife_print_gpu.json').write_text(json.dumps(report,indent=2))

for n,name in NAMES.items():
    shutil.copy2(OUT/(name+'.pptx'),TARGET/(name+'.pptx'))
    pdf=ROOT/'qa/renders/chemicals_of_life'/f'l{n}'/(name+'.pdf')
    if pdf.exists():shutil.copy2(pdf,TARGET/(name+'.pdf'))
    for suffix in ['manifest','validation','gpu']:
        p=ROOT/'qa'/f'{name}_{suffix}.json'
        if p.exists():
            (TARGET/'qa').mkdir(exist_ok=True);shutil.copy2(p,TARGET/'qa'/p.name)
    log=ROOT/'build_logs'/f'{name}.md'
    if log.exists():
        (TARGET/'build_logs').mkdir(exist_ok=True);shutil.copy2(log,TARGET/'build_logs'/log.name)
for p in [ROOT/'qa/ChemicalsOfLife_sources.json',ROOT/'qa/ChemicalsOfLife_checks.json',ROOT/'qa/ChemicalsOfLife_print_gpu.json',ROOT/'assets/chemicals_of_life/provenance.json']:
    (TARGET/'qa').mkdir(exist_ok=True);shutil.copy2(p,TARGET/'qa'/p.name)
shutil.copy2(ROOT/'content/chemicals_of_life/README.md',TARGET/'README.md')
files=sorted(p for p in TARGET.rglob('*') if p.is_file() and p.name!='package_manifest.json')
manifest={'status':'REVIEW DRAFT — NOT APPROVED','files':[{'path':str(p.relative_to(TARGET)),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in files]}
(TARGET/'package_manifest.json').write_text(json.dumps(manifest,indent=2))
archive=ROOT/'outputs/drafts/ChemicalsOfLife_Lessons1-6_Review.zip'
with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED) as z:
    for p in sorted(TARGET.rglob('*')):
        if p.is_file():z.write(p,Path(TARGET.name)/p.relative_to(TARGET))
with zipfile.ZipFile(archive) as z:
    assert z.testzip() is None
print(archive,archive.stat().st_size,hashlib.sha256(archive.read_bytes()).hexdigest())
