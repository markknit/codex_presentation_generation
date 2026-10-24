"""Reusable editable PowerPoint visual grammar. No lesson artwork is copied from references."""
from io import BytesIO
from lxml import etree
from pathlib import Path
import json,re,hashlib
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE as M
from pptx.enum.text import PP_ALIGN as A
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from build_biology_cellstructure_l2_l3 import shape,text,line,rgb,note,table,COL,DQ,ROOT,DATA
INK='172C38';MUTED='52636A';CREAM='FFF9EB';PALE={'Predict':'E9F8FC','Observe':'F7F9DA','Explain':'FBE9F5','Driving Question Board':'FFF2DC','Model Building':'FDEBED'}
ASSETS=ROOT/'assets/cellstructure_v2'

def tx(s,t,x,y,w,h,size=23,color=INK,bold=False,align=A.LEFT):return text(s,t,x,y,w,h,size,color,bold,align)
def box(s,x,y,w,h,fill='F0F4F4',border=None):return shape(s,x,y,w,h,fill,border)
def oval(s,x,y,w,h,fill,border=None):return shape(s,x,y,w,h,fill,border,M.OVAL)
def badge(s,t,x,y,c,size=.44):
 oval(s,x,y,size,size,c);tx(s,t,x,y+.06,size,size-.06,15,'FFFFFF' if c==INK else INK,True,A.CENTER)
def timer(s,t,label='silent work',x=10.65,y=5.95):
 box(s,x,y,2.0,.69,INK);tx(s,t,x+.1,y+.04,1.8,.37,25,'FFFFFF',True,A.CENTER);tx(s,label,x+.08,y+.44,1.84,.2,10,'FFFFFF',False,A.CENTER)
def arrow(s,x,y,w=.42,c='CF128C'):
 shape(s,x,y,w,.26,c,kind=M.CHEVRON)
def question(s,t,phase='Predict',y=2.28,w=11.9,size=31):
 box(s,.65,y,w,1.5,PALE.get(phase,'F0F4F4'),COL.get(phase,'CBD5D8'));badge(s,'?',.88,y+.45,COL.get(phase,'47C6E4'),.55);tx(s,t,1.65,y+.24,w-1.22,1.12,size,INK,True)
def steps(s,items,x=.75,y=2.35,w=11.8,dy=1.0,size=25,c='47C6E4'):
 for i,t in enumerate(items):badge(s,str(i+1),x,y+i*dy,c,.48);tx(s,t,x+.75,y+i*dy+.02,w-.75,dy-.12,size)
def photo(s,name,x=8.25,y=2.08,w=4.25):
 p=ASSETS/(name+'.png')
 if not p.exists():raise FileNotFoundError(p)
 s.shapes.add_picture(str(p),Inches(x),Inches(y),width=Inches(w),height=Inches(w))
def board(s,x=8.5,y=2.4,w=3.9):
 box(s,x,y,w,3.3,'D6AB73');box(s,x+.12,y+.12,w-.24,3.06,'FFF4DA')
 for i in range(6):
  xx=x+.35+i%3*1.13;yy=y+.4+i//3*1.35
  box(s,xx,yy,.92,1.05,['FFF09B','CBE6B1','B8E8F4'][i%3]);oval(s,xx+.4,yy+.06,.12,.12,'E78165')
 tx(s,'Our ideas keep growing',x,y+3.5,w,.4,17,MUTED,False,A.CENTER)
def notebook(s,x=8.5,y=2.5,w=3.8):
 box(s,x+.08,y+.08,w,3.1,'CAD9DF');box(s,x,y,w,3.1,CREAM,'D4C7A9');line(s,x+.38,y+.2,x+.38,y+2.9,'E78484')
 for i in range(6):line(s,x+.55,y+.55+i*.4,x+w-.22,y+.55+i*.4,'BCD9E2')
 tx(s,'YOUR MODEL',x+.65,y+.17,w-.9,.3,14,MUTED,True)
 shape(s,x+w-.8,y+2.3,.22,1.1,'F7C759',kind=M.RECTANGLE).rotation=35

def organelle(s,kind,x,y,w,h):
 # The same editable primitives scale from a labelled whole cell to an organelle detail.
 if kind=='chloroplast':
  oval(s,x,y,w,h,'87BA74','347954');oval(s,x+w*.04,y+h*.08,w*.92,h*.84,'C6E4A3','699D61')
  for i in range(3):
   if i<2:line(s,x+w*(.28+i*.25),y+h*.5,x+w*(.53+i*.25),y+h*.5,'43885B',max(.6,w*.35))
   for j in range(4):box(s,x+w*(.18+i*.25),y+h*(.22+j*.145),w*.16,h*.09,'43885B')
 elif kind=='mitochondrion':
  oval(s,x,y,w,h,'F0B574','BC7C43');oval(s,x+w*.045,y+h*.10,w*.91,h*.8,'FFE0AA','C28D55')
  pts=[(.15,.22),(.25,.22),(.25,.67),(.32,.67),(.32,.18),(.46,.18),(.46,.73),(.53,.73),(.53,.18),(.67,.18),(.67,.67),(.74,.67),(.74,.22),(.85,.22)]
  for a,b in zip(pts,pts[1:]):line(s,x+w*a[0],y+h*a[1],x+w*b[0],y+h*b[1],'AD733F',max(.6,w*.45))
 elif kind=='nucleus':
  oval(s,x,y,w,h,'B298D1','7D5FA6');oval(s,x+w*.04,y+h*.06,w*.92,h*.88,'DECEF1','AA90C3')
  oval(s,x+w*.37,y+h*.36,w*.25,h*.25,'7D5FA6')
  for xx,yy in [(.2,.3),(.62,.25),(.19,.7),(.7,.65)]:
   line(s,x+w*xx,y+h*yy,x+w*(xx+.11),y+h*(yy-.04),'A489C4',max(.5,w*.3));line(s,x+w*(xx+.11),y+h*(yy-.04),x+w*(xx+.15),y+h*(yy+.09),'A489C4',max(.5,w*.3))
 elif kind=='vacuole':
  box(s,x,y,w,h,'BDEAF1','64AFC0');line(s,x+w*.12,y+h*.15,x+w*.87,y+h*.15,'EAFBFD',max(.6,w*.6));line(s,x+w*.12,y+h*.19,x+w*.12,y+h*.8,'EAFBFD',max(.6,w*.6))
 elif kind in ['rough ER','Golgi apparatus']:
  for i in range(4):
   xx=x+w*(.04+i*.045);yy=y+h*(.10+i*.21);ww=w*(.85-i*.09)
   box(s,xx,yy,ww,h*.13,'C4A2D3' if kind=='rough ER' else 'ED9D8C')
   line(s,xx+ww*.1,yy+h*.04,xx+ww*.9,yy+h*.04,'EAD8EF' if kind=='rough ER' else 'F9C4B5',max(.5,w*.25))
   if kind=='rough ER':
    for j in range(5):oval(s,xx+ww*(.06+j*.18),yy-h*.025,max(.03,w*.025),max(.03,w*.025),INK)
    if i<3:line(s,xx+ww*.12,yy+h*.13,xx+ww*.12,yy+h*.22,'C4A2D3',max(.7,w*.8))
   else:oval(s,x+w*(.87-i*.025),yy+h*.03,w*.07,h*.11,'F6B3A4','D78577')
 elif kind=='ribosomes':
  for i in range(6):
   xx=x+w*(.14+(i%3)*.3);yy=y+h*(.25+(i//3)*.36)
   oval(s,xx,yy,w*.10,h*.13,'4B6274');oval(s,xx+w*.015,yy+h*.08,w*.075,h*.06,'6B879A')

def cell(s,kind='plant',x=4.25,y=2.45,w=4.8,h=3.75,labels=True,focus=None):
 """Illustrative generalised section; same palette and organelle geometry across lessons."""
 points={}
 if kind=='bacterium':
  a=box(s,x,y+.6,w,h-1.2,'D5B783','876239');a.adjustments[0]=.5
  a=box(s,x+.09,y+.69,w-.18,h-1.38,'FFF8D5','B64475');a.adjustments[0]=.5
  pts=[(.26,.43),(.44,.35),(.64,.55),(.38,.62),(.64,.36)]
  for i in range(16):oval(s,x+.35+(i*13%37)/10,y+.95+(i*7%19)/10,.06,.06,INK)
  for a,b in zip(pts,pts[1:]+pts[:1]):line(s,x+w*a[0],y+h*a[1],x+w*b[0],y+h*b[1],'398EAC',2)
  points={'Cell wall':(x+w*.28,y+.6),'Cell membrane':(x+.11,y+.6+(h-1.2)/2),'Cytoplasm':(x+w*.65,y+h*.32),'DNA in nucleoid':(x+2.2,y+2.1),'Ribosomes':(x+3.6,y+2.7)}
 else:
  if kind=='plant':
   box(s,x,y,w,h,'8CBB7A','397047');box(s,x+.13,y+.13,w-.26,h-.26,'EEF6D5','B43B77')
   organelle(s,'vacuole',x+w*.33,y+h*.24,w*.48,h*.58)
   for xx,yy in [(.17,.15),(.17,.78),(.69,.13)]:organelle(s,'chloroplast',x+w*xx,y+h*yy,w*.19,h*.12)
   points={'Cell wall':(x,y+.6),'Cell membrane':(x+.13,y+1.2),'Cytoplasm':(x+w*.48,y+h*.89),'Large central vacuole':(x+w*.57,y+h*.53),'Chloroplast':(x+w*.78,y+h*.19)}
  else:
   oval(s,x,y,w,h,'F8E9EB','B43B77');points={'Cell membrane':(x+.15,y+h*.35),'Cytoplasm':(x+w*.62,y+h*.65)}
  organelle(s,'nucleus',x+w*.08,y+h*.32,w*.22,h*.25)
  organelle(s,'mitochondrion',x+w*.6,y+h*(.85 if kind=='plant' else .80),w*.21,h*.10)
  organelle(s,'rough ER',x+w*.10,y+h*.61,w*.25,h*.13)
  organelle(s,'Golgi apparatus',x+w*(.65 if kind=='plant' else .5),y+h*(.03 if kind=='plant' else .12),w*.22,h*.10)
  for xx,yy in [(.23,.86),(.30,.21),(.87,.4),(.87,.65)]:oval(s,x+w*xx,y+h*yy,.065,.065,INK)
  points.update({'Nucleus':(x+w*.19,y+h*.445),'Mitochondrion':(x+w*.70,y+h*(.90 if kind=='plant' else .85)),'Rough ER':(x+w*.23,y+h*.67),'Ribosomes':(x+w*.87,y+h*.65),'Golgi apparatus':(x+w*(.76 if kind=='plant' else .61),y+h*(.08 if kind=='plant' else .17))})
 if focus:
  px,py=points[focus];oval(s,px-.16,py-.16,.32,.32,'FFFFFF','CF128C')
 if labels:
  keys=list(points)
  left=([k for k in ['Cell wall','Cell membrane','Nucleus','Rough ER','Cytoplasm'] if k in points] if kind!='bacterium' else ['Cell wall','Cell membrane','DNA in nucleoid']);right=[k for k in ['Golgi apparatus','Chloroplast','Large central vacuole','Ribosomes','Mitochondrion'] if k in points] if kind!='bacterium' else ['Cytoplasm','Ribosomes']
  for side,ks in enumerate([left,right]):
   for i,k in enumerate(ks):
    yy=2.65+i*.7;xx=.65 if side==0 else 9.7;ww=3.0 if side==0 else 2.95
    tx(s,k,xx,yy,ww,.63,17,INK,True,A.RIGHT if side==0 else A.LEFT)
    px,py=points[k];endx=3.8 if side==0 else 9.4
    line(s,px,py,endx,yy+.28,'7D8D94',.9);oval(s,px-.035,py-.035,.07,.07,INK)
 tx(s,'Editable drawing · not a micrograph · not to scale',.65,6.57,10.0,.26,12,MUTED)
 return points

def kitchen(s):
 photo(s,'kitchen',7.8,2.35,4.25)

def onion(s,x=8.7,y=3.0):
 oval(s,x,y,2.5,2.35,'DDA7AB','8F5E70')
 for i in range(4):oval(s,x+.2+i*.16,y+.2+i*.16,2.1-i*.32,1.95-i*.32,'F9E4D3','B97887')
 for i in range(4):line(s,x+1.1,y+.1,x+.5+i*.45,y-.5,'669864',2)
 tx(s,'Inner bulb layers',x-.2,y+2.6,3,.4,18,MUTED,False,A.CENTER)

def wetmount(s):
 for i,title in enumerate(['Peel a thin layer','Lay flat + stain','Lower the coverslip']):
  xx=.7+i*4.18;badge(s,str(i+1),xx,2.3,'D7DF23');tx(s,title,xx+.65,2.3,3.3,.65,22,INK,True)
  box(s,xx+.3,4.3,3.05,.24,'C8EDF0','4794A3');box(s,xx+1.2,4.05,1.0,.17,'ECD9A8','C5A169')
  if i==0:
   oval(s,xx+.75,3.15,1.8,.8,'E5BBC0','AB7886');line(s,xx+1.7,3.8,xx+2.5,3.0,'B17A8B',2)
  if i==1:
   shape(s,xx+1.6,3.35,.25,.4,'AF864F',kind=M.OVAL)
  if i==2:
   ln=line(s,xx+1.3,4.04,xx+2.5,3.28,'4794A3',3)
 tx(s,'One drop of the teacher-approved stain · follow the demonstration',.75,5.6,11.8,.6,24,INK,True)

def microscope(s,x=7.9,y=2.25):
 photo(s,'microscope',x+.05,y+.2,4.25)

class VisualDeck:
 def __init__(self,n,name,*,subject='Biology',grade=10,unit='Cell Structure',template=None):
  self.n=n;self.name=name;self.subject=subject;self.grade=grade;self.unit=unit;self.records=[]
  self.deck=Presentation(template or ROOT/'references/templates/Bio_G10_CellStructure_L2_What_Is_a_Cell.pptx')
  # Resolve the notes master's actual theme; custom templates may use other part names.
  for master in self.deck.part.package.iter_parts():
   if master.content_type=='application/vnd.openxmlformats-officedocument.presentationml.notesMaster+xml':
    for rel in master.rels.values():
     if rel.reltype.endswith('/theme'):
      part=rel.target_part;e=etree.fromstring(part.blob)
      for node in e.xpath('//*[local-name()="majorFont" or local-name()="minorFont"]/*[local-name()="latin"]'):node.set('typeface','Calibri')
      part._blob=etree.tostring(e,xml_declaration=True,encoding='UTF-8',standalone=True)
  self.logos=[(sh.image.blob,sh.left,sh.top,sh.width,sh.height) for sh in self.deck.slides[0].shapes if sh.shape_type==13]
  for sid in list(self.deck.slides._sldIdLst):self.deck.part.drop_rel(sid.rId);self.deck.slides._sldIdLst.remove(sid)
 def slide(self,title,phase,notes,sub='',hidden=False,dark=False,check=False):
  s=self.deck.slides.add_slide(self.deck.slide_layouts[0]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(INK if dark else 'FFFFFF')
  if hidden:s._element.set('show','0')
  c=COL.get(phase,'E44145' if phase=='Safety' else '47C6E4');box(s,0,0,13.333,.075,c)
  lab='TEACHER ONLY · HIDDEN' if hidden else phase.upper();wide=3.3 if hidden or phase=='Driving Question Board' else 1.8
  box(s,.65,.36,wide,.35,c if not hidden else INK);tx(s,lab,.75,.43,wide-.2,.2,10,'FFFFFF' if hidden or phase in ['Explain','Model Building'] else INK,True,A.CENTER)
  if title:tx(s,title,.65,.98,12,1.0,32,'FFFFFF' if dark else INK,True)
  if sub:tx(s,sub,.68,1.84,11.9,.42,18,'D2E4E8' if dark else MUTED)
  if check:tx(s,'CHECK FOR UNDERSTANDING',9.4,.43,3.25,.22,10,'FFFFFF' if dark else INK,True,A.RIGHT)
  tx(s,f'© 2026 SeaVuria & ARES Education · CC BY-NC 4.0 · {self.subject} G{self.grade} · {self.unit} · Lesson {self.n}',.65,7.04,9.1,.22,9,'D2E4E8' if dark else MUTED)
  if dark:box(s,10.02,6.82,2.65,.52,'FFFFFF')
  for blob,l,t,w,h in self.logos:s.shapes.add_picture(BytesIO(blob),l,t,w,h)
  if s.notes_slide.notes_text_frame is None:
   xml='<p:sp '+nsdecls('p','a')+'><p:nvSpPr><p:cNvPr id="99" name="Notes Text"/><p:cNvSpPr txBox="1"/><p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>'
   s.notes_slide.shapes._spTree.insert_element_before(parse_xml(xml),'p:extLst')
  s.notes_slide.notes_text_frame.text=notes
  self.records.append({'slide':len(self.records)+1,'title':title,'phase':phase,'minutes':float(re.search(r'TIME: ([\d.]+)',notes)[1]),'hidden':hidden})
  return s
 def save(self,out_dir=None):
  folder=Path(out_dir) if out_dir else ROOT/'outputs/drafts/v2';folder.mkdir(parents=True,exist_ok=True)
  p=folder/f'{self.name}.pptx'
  for slide in self.deck.slides:
   for para in slide.notes_slide.notes_text_frame.paragraphs:
    para.font.name='Calibri'
    for run in para.runs:run.font.name='Calibri'
  self.deck.save(p)
  (ROOT/'qa'/f'{self.name}_manifest.json').write_text(json.dumps({'deck':str(p),'slides':self.records,'total_minutes':sum(r['minutes'] for r in self.records),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'status':'REVIEW DRAFT — NOT APPROVED'},indent=2))
  print(p)
