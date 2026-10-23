#!/usr/bin/env python3
"""Structural checks for the L1 review pilot; human release gates remain separate."""
import hashlib
import json
import re
import subprocess
import zipfile
from pathlib import Path
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
NAME='Bio_G10_CellStructure_L1_The_Healing_Cut'
path=ROOT/'outputs/drafts'/f'{NAME}.pptx'
deck=Presentation(path)
errors=[];times=[];hidden=[];checks=[]
labels=['PURPOSE OF THIS LESSON','BEFORE CLASS','DO','SAY','WAIT TIME','LOOK FOR','IF…','WHY THIS WORKS','ANSWER KEY','SCIENCE NOTE FOR YOU','TEACHER REFLECTION']
for i,s in enumerate(deck.slides,1):
    if s._element.get('show')=='0':hidden.append(i)
    note=s.notes_slide.notes_text_frame.text
    found=re.findall(r'^TIME: ([\d.]+) minutes$',note,re.M)
    if len(found)!=1:errors.append(f'slide {i}: expected one numeric TIME')
    else:times.append(float(found[0]))
    positions=[note.index('\n'+x+'\n') for x in labels if '\n'+x+'\n' in note]
    if positions!=sorted(positions):errors.append(f'slide {i}: note label order')
    alltext='\n'.join(sh.text for sh in s.shapes if sh.has_text_frame)
    if 'CHECK FOR UNDERSTANDING' in alltext:
        checks.append(i)
        for x in ['LOOK FOR','IF…']:
            if not re.search('^'+re.escape(x)+'$',note,re.M):errors.append(f'slide {i}: check lacks {x}')
    if '© 2026 SeaVuria & ARES Education' not in alltext:errors.append(f'slide {i}: footer missing')
    for sh in s.shapes:
        if sh.width>0 and sh.height>0 and (sh.left<-20000 or sh.top<-20000 or sh.left+sh.width>deck.slide_width+20000 or sh.top+sh.height>deck.slide_height+20000):errors.append(f'slide {i}: shape outside canvas {sh.name}')
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.name!='Calibri':errors.append(f'slide {i}: non-Calibri run')
    if i not in [2,3] and re.search(r'\bDQB\b',alltext):errors.append(f'slide {i}: abbreviated Driving Question Board')
if hidden!=[2,3]:errors.append('expected only slides 2 and 3 hidden')
if sum(times)!=40:errors.append(f'total timing {sum(times)} != 40')
if len(deck.slides)!=21:errors.append('expected 21 slides')
if abs(deck.slide_width/deck.slide_height-16/9)>.01:errors.append('not 16:9')
with zipfile.ZipFile(path) as z:
    bad=z.testzip()
    if bad:errors.append('bad ZIP member '+bad)
report={'deck':str(path.relative_to(ROOT)),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'structural_status':'PASS' if not errors else 'FAIL','errors':errors,'slides':len(deck.slides),'hidden_slides':hidden,'times_minutes':times,'total_minutes':sum(times),'formative_check_slides':checks,'render_font':subprocess.check_output(['fc-match','Calibri'],text=True).strip(),'release_status':'DRAFT — NOT APPROVED','pending':['Teacher source/science review and all build-log decisions','Calibri fidelity on review machine','Native PowerPoint Slide Show and Presenter View','ARES link/content/offline verification','Manual versus source launch-lesson exceptions'], 'limits':'Shape bounds are not a guarantee of text fit; visual inspection is recorded separately. CUDA pixel signals are not semantic or science review.'}
(ROOT/'qa'/f'{NAME}_validation.json').write_text(json.dumps(report,indent=2))
print(json.dumps(report,indent=2))
raise SystemExit(bool(errors))
