#!/usr/bin/env python3
"""Build editable review drafts for Cell Structure L2 and L3, explicitly authorised on 2026-10-23."""
from pathlib import Path
from io import BytesIO
import hashlib,json,math
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
ROOT=Path(__file__).resolve().parents[1]
UNIT=ROOT/'inputs/v2/Biology/SS1.1_Cell_Structure'
DATA=json.loads((UNIT/'Biology_Cell_Structure_data.json').read_text())
COL={'Predict':'47C6E4','Observe':'D7DF23','Explain':'CF128C','Driving Question Board':'FCBA63','Model Building':'E44145'}
DQ='Why do cells in our body look so different from each other — and how does the shape of a cell determine the job it can do?'
NAMES={2:'Bio_G10_CellStructure_L2_What_Is_a_Cell',3:'Bio_G10_CellStructure_L3_Prokaryotic_vs_Eukaryotic_Cells'}
TEMPLATE=ROOT/'references/templates/Bio_G10_CellStructure_L2_What_Is_a_Cell.pptx'
SOURCES=[str(p.relative_to(ROOT)) for p in sorted(UNIT.glob('*'))]+[str(TEMPLATE.relative_to(ROOT)),'references/instructions/SeaVuria_Science_Lesson_PPT_Build_Manual.docx']

def rgb(s):return RGBColor.from_string(s)

def shape(s,x,y,w,h,fill='F4F5F8',line=None,kind=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if line:sh.line.color.rgb=rgb(line)
    else:sh.line.fill.background()
    if kind==MSO_SHAPE.ROUNDED_RECTANGLE:
        try:sh.adjustments[0]=.08
        except (IndexError,ValueError):pass
    return sh

def text(s,t,x,y,w,h,size=24,color='161616',bold=False,align=PP_ALIGN.LEFT):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
    tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align
        p.space_after=Pt(5);p.font.name='Calibri';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
        r=p.add_run();r.text=line;r.font.name='Calibri';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
    return sh

def line(s,x1,y1,x2,y2,color='687386',width=1):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(width);return sh

def card(s,title,body,x,y,w,h,fill='F4F5F8',size=23):
    shape(s,x,y,w,h,fill);text(s,title,x+.22,y+.18,w-.44,.6,19,bold=True);text(s,body,x+.22,y+.92,w-.44,h-1.04,size)

def timer(s,label,x=10.7,y=5.95):
    shape(s,x,y,1.9,.55,'161616');text(s,label,x+.1,y+.09,1.7,.4,20,'FFFFFF',True,PP_ALIGN.CENTER)

def table(s,headers,rows,x=.65,y=2.05,w=12.05,h=3.9,widths=None,size=18):
    sh=s.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(w),Inches(h));t=sh.table
    if widths:
        for c,v in zip(t.columns,widths):c.width=Inches(v)
    for r,vals in enumerate([headers]+rows):
        for c,v in enumerate(vals):
            cell=t.cell(r,c);cell.text=v;cell.fill.solid();cell.fill.fore_color.rgb=rgb('161616' if r==0 else ('F4F5F8' if r%2 else 'FFFFFF'))
            cell.margin_left=cell.margin_right=Inches(.12);cell.margin_top=Inches(.06);cell.margin_bottom=Inches(.04);cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                for run in p.runs:run.font.name='Calibri';run.font.size=Pt(size);run.font.bold=r==0;run.font.color.rgb=rgb('FFFFFF' if r==0 else '161616')
    return sh

def note(minutes,do,say='',wait='',look='',iff='',why='',key='',science='',reflection='',purpose='',before='',phase_note=''):
    d=[('PURPOSE OF THIS LESSON',purpose),('BEFORE CLASS',before),('DO',do),('SAY',say),('WAIT TIME',wait),('LOOK FOR',look),('IF…',iff),('WHY THIS WORKS',why),('ANSWER KEY',key),('SCIENCE NOTE FOR YOU',science),('TEACHER REFLECTION',reflection)]
    return (phase_note+'\n' if phase_note else '')+f'TIME: {minutes:g} minutes\n'+''.join('\n'+k+'\n'+v+'\n' for k,v in d if v)

class Builder:
    def __init__(self,n):
        self.n=n;self.lesson=DATA['LESSONS'][n-1];self.deck=Presentation(TEMPLATE);self.records=[]
        self.logos=[(sh.image.blob,sh.left,sh.top,sh.width,sh.height) for sh in self.deck.slides[0].shapes if sh.shape_type==13]
        for sid in list(self.deck.slides._sldIdLst):self.deck.part.drop_rel(sid.rId);self.deck.slides._sldIdLst.remove(sid)
    def slide(self,phase,title,minutes,do,*,hidden=False,dark=False,check=False,**kw):
        s=self.deck.slides.add_slide(self.deck.slide_layouts[0]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('111111' if dark else 'FFFFFF')
        if hidden:s._element.set('show','0')
        lab='TEACHER ONLY · HIDDEN' if hidden else phase.upper();fill='161616' if hidden else COL.get(phase,'161616')
        wide=3.6 if phase=='Driving Question Board' or hidden else 2.0
        shape(s,.6,.45,wide,.38,fill);text(s,lab,.68,.51,wide-.16,.25,10,'FFFFFF' if hidden or phase not in COL or phase in ['Explain','Model Building'] else '161616',True,PP_ALIGN.CENTER)
        if title:text(s,title,.6,1.03,12.1,1.0,32,'FFFFFF' if dark else '161616',True)
        if check:text(s,'✓ CHECK FOR UNDERSTANDING',9.0,.5,3.7,.3,11,'FFFFFF' if dark else '333333',True,PP_ALIGN.RIGHT)
        text(s,f'© 2026 SeaVuria & ARES Education · CC BY-NC 4.0 · Biology G10 · Cell Structure · Lesson {self.n}',.6,7.03,8.7,.25,9,'B8BCC5' if dark else '697080')
        shape(s,9.98,6.84,2.75,.5,'FFFFFF','E1E4EA')
        for blob,l,t,w,h in self.logos:s.shapes.add_picture(BytesIO(blob),l,t,w,h)
        kw['science']=kw.get('science','')+'\nSource: V2 Cell Structure, Lesson '+str(self.n)+'. Review draft: native PowerPoint, ARES and teacher sign-off remain unverified.'
        if s.notes_slide.notes_text_frame is None:
            xml='<p:sp '+nsdecls('p','a')+'><p:nvSpPr><p:cNvPr id="99" name="Notes Text"/><p:cNvSpPr txBox="1"/><p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>'
            s.notes_slide.shapes._spTree.insert_element_before(parse_xml(xml),'p:extLst')
        s.notes_slide.notes_text_frame.text=note(minutes,do,**kw)
        self.records.append({'slide':len(self.deck.slides),'phase':phase,'minutes':minutes,'title':title,'check':check,'hidden':hidden})
        return s
    def opening(self,short,desc,goals,words,phases):
        n=self.n
        s=self.slide('Predict','',.25,'Introduce the lesson title briefly. The opening is included in the Predict allocation.',dark=True,purpose=self.lesson['slo']['purposeInStoryline'],before=('Prepare the class Driving Question Board, notebooks, coloured pencils, timer and paper slips. '+('Provide 12 pre-cut feature cards per group of four, diagram sheets and the three-scientist reading/evidence sheet. Have Version 0 models ready.' if n==2 else 'Provide the E. coli and cheek-cell diagrams and five-row comparison table. Have Version 1 models ready. No live cultures are used.')),say=desc)
        text(s,f'GRADE 10 BIOLOGY · CELL STRUCTURE · LESSON {n} OF 12 · {80 if n==2 else 40} MINUTES',.8,1.1,11.8,.4,14,'D7DF23',True)
        text(s,short,.8,2,8.2,1.7,46,'FFFFFF',True);text(s,desc,.8,4,7.8,1.5,22,'D1D3D9')
        for i,(ph,c) in enumerate(COL.items()):
            shape(s,9.2,2.0+i*.65,.35,.35,c,kind=MSO_SHAPE.OVAL);text(s,ph,9.75,2+i*.65,3,.6,17,'FFFFFF',True)
        text(s,'REVIEW DRAFT · 23 OCTOBER 2026',.8,6.2,8,.3,11,'FCBA63')
        s=self.slide('Routine','How these lessons work',0,'Read this hidden orientation before class. Keep slides 2 and 3 hidden. Use Presenter View to read the teacher notes.',hidden=True,science='Notes labels follow the approved manual: TIME, DO, SAY, WAIT TIME, LOOK FOR, IF…, WHY THIS WORKS, ANSWER KEY, SCIENCE NOTE FOR YOU, TEACHER REFLECTION. Timing belongs to the visible slides only.')
        for i,(ph,c) in enumerate(COL.items()):
            shape(s,.65,2.1+i*.78,2.7,.55,c);text(s,ph,.73,2.2+i*.78,2.54,.4,14,'FFFFFF' if ph in ['Explain','Model Building'] else '161616',True,PP_ALIGN.CENTER)
            body=['Write independently before hearing others.','Record evidence from the activity.','Use evidence to revise and explain.','Track answers and new questions.','Revise the same model as thinking grows.'][i]
            text(s,body,3.6,2.13+i*.78,5.0,.65,17)
        card(s,'Three habits','Wait after questions.\nRecord predictions without correction.\nChoose pupils across the class.',9.1,2.1,3.5,3.9,size=20)
        s=self.slide('Glance',f'Lesson {n} at a glance',0,'Prepare before pupils arrive. Phase timings include the opening and closing work. Materials and notes support offline teaching.',hidden=True,science='All resource buttons copy V2 ARES URLs, untested locally. Preview first: comparison material can fit after Predict; detailed organelle and endosymbiosis resources must not reveal later lessons.',iff='No projector: use printed diagrams and reading/table sheets. No sticky notes: paper slips and tape. Pre-cut feature cards before L2; no pupil cutting task is required.')
        rows=[[p,str(t)+' min',act] for p,t,act in phases]
        table(s,['Phase','Time','Pupil activity'],rows,.65,2.0,8.1,4.15,[2.3,.9,4.9],15)
        resources=[];seen=set()
        for phase in ['predict','observe','model']:
            for item in self.lesson['resourceLinks'][phase].values():
                if isinstance(item,dict) and item.get('direct_url') and item['direct_url'] not in seen:
                    resources.append(item);seen.add(item['direct_url'])
        resource_lines=[];all_seen=set()
        for phase_key,phase_items in self.lesson['resourceLinks'].items():
            for item in phase_items.values():
                if isinstance(item,dict) and item.get('direct_url') and item['direct_url'] not in all_seen:
                    title=item['title'];all_seen.add(item['direct_url'])
                    decision='Do not use today; defer detailed organelles/endosymbiosis to the appropriate later preparation.' if any(word in title.lower() for word in ['endosymbiosis','organelles','plant cell walls']) else 'Preview first; not an exact match established. Use only after the Predict phase and check for later-content spoilers.'
                    resource_lines.append(title+' — '+decision+' Source URL: '+item['direct_url'])
        s.notes_slide.notes_text_frame.text+='\n\nResource match audit (all source resources; links/content not verified):\n'+'\n'.join(resource_lines)
        text(s,'ARES · PREVIEW FIRST',9.1,2.05,3.6,.35,12,bold=True)
        for i,item in enumerate(resources[:3]):
            b=shape(s,9.1,2.65+i*.85,3.55,.65,'161616');b.click_action.hyperlink.address=item['direct_url']
            sh=text(s,item['title'],9.27,2.76+i*.85,3.2,.5,13,'FFFFFF',True);sh.click_action.hyperlink.address=item['direct_url']
        text(s,'Printed diagrams and reading sheets are the offline fallback.\n\nResources are not verified.',9.1,5.3,3.5,1.1,16)
        s=self.slide('Unit map','Our place in the investigation',.5 if n==2 else .25,'Briefly link to the previous lesson and point to today. Keep later detailed mechanisms for the appropriate source lesson.',say='Last lesson we recorded '+('our starting ideas about Amina’s cut.' if n==2 else 'cell theory and our first labelled model.'))
        titles=['The healing cut','What is a cell?','Prokaryotes and eukaryotes','Plant cells','Animal cells','Onion practical','Cheek practical','Animal specialisation','Plant specialisation','Organisation','Synthesis','Final explanation']
        for i,title in enumerate(titles):
            col=i%6;row=i//6;x=.65+col*2.05;y=2.2+row*1.1
            shape(s,x,y,.45,.45,'47C6E4' if i+1==n else 'EEF0F4',kind=MSO_SHAPE.OVAL);text(s,str(i+1),x,y+.06,.45,.3,16,bold=True,align=PP_ALIGN.CENTER);text(s,title,x+.55,y,1.45,.85,14,bold=i+1==n)
        card(s,'Driving question',DQ,.65,4.7,12.05,1.7,'E4F7FC',20)
        s=self.slide('Today','By the end of this lesson you can…',.75 if n==2 else .5,'Introduce the three learning goals and process vocabulary. Do not explain the comparison before pupils predict.',say='We will use evidence to test our ideas and connect them to Amina’s healing finger.')
        for i,g in enumerate(goals):
            shape(s,.65,2.1+i*.95,.5,.5,list(COL.values())[i],kind=MSO_SHAPE.OVAL);text(s,g,1.4,2.12+i*.95,7.0,.8,23,bold=True)
        card(s,'Key words',words,9.1,2.1,3.5,3.8,size=18)
    def save(self,total):
        assert sum(r['minutes'] for r in self.records)==total,(self.n,sum(r['minutes'] for r in self.records))
        self.deck.core_properties.title=DATA['LESSONS'][self.n-1]['title'];self.deck.core_properties.subject='Biology Grade 10 • Cell Structure • review draft';self.deck.core_properties.comments='User authorised draft continuation beyond pilot checkpoint. Not approved for release.'
        target=ROOT/'outputs/drafts'/f'{NAMES[self.n]}.pptx';self.deck.save(target)
        manifest={'lesson':self.n,'status':'DRAFT — NOT APPROVED','authorisation':'User: Proceed with lessons 2-3; 2026-10-23. Overrides pre-generation pilot wait for these drafts only; release gates remain open.','total_minutes':total,'slides':self.records,'inputs':[{'path':p,'sha256':hashlib.sha256((ROOT/p).read_bytes()).hexdigest()} for p in SOURCES],'output_sha256':hashlib.sha256(target.read_bytes()).hexdigest()}
        (ROOT/'qa'/f'{NAMES[self.n]}_manifest.json').write_text(json.dumps(manifest,indent=2));print(target)

# Every organelle/label is a native PowerPoint shape/text item, not a bitmap.
def diagram(s,kind,x,y,w,h,labels=True):
    # Body area left, numbered label list right; all numbers are cross-referenced.
    bw=w*.57;bx=x+.12;bh=min(h-.65,bw*.5) if kind=='bacterium' else h-.65;by=y+(h-bh)/2
    plant=kind=='plant';bact=kind=='bacterium'
    shape(s,bx,by,bw,bh,'E4F1DB' if plant else ('F1F0FC' if bact else 'FBEAF2'),'657452' if plant else '606B83',MSO_SHAPE.ROUNDED_RECTANGLE if plant or bact else MSO_SHAPE.OVAL)
    shape(s,bx+.09,by+.09,bw-.18,bh-.18,'E7F8FD' if plant else ('F6F4FC' if bact else 'FBEAF2'),'4B6B8A',MSO_SHAPE.ROUNDED_RECTANGLE if plant or bact else MSO_SHAPE.OVAL)
    points=[]
    if bact:
        # Closed, folded chromosome; the nucleoid is not enclosed in a membrane.
        for k in range(25):
            angle=2*math.pi*k/24
            px=bx+bw*(.5+.28*math.cos(angle));py=by+bh*(.5+.13*math.sin(angle)+.04*math.sin(3*angle))
            if k:line(s,prev[0],prev[1],px,py,'9670BF',2)
            prev=(px,py)
        points=[(bx+.02,by+.25),(bx+.13,by+bh-.25),(bx+bw*.7,by+bh*.3),(bx+bw*.5,by+bh*.5),(bx+bw*.2,by+bh*.75)]
        names=['Cell wall (peptidoglycan)','Cell membrane','Cytoplasm','Circular DNA in nucleoid','Ribosomes']
        # Flagellum leaves body at right, visibly separate from numbered list.
        for k in range(5):
            line(s,bx+bw+k*.055,by+bh*.66+math.sin(k*1.4)*.08,bx+bw+(k+1)*.055,by+bh*.66+math.sin((k+1)*1.4)*.08,'606B83',1.5)
        names+=['Flagellum'];points.append((bx+bw,by+bh*.66))
    else:
        if plant:shape(s,bx+bw*.42,by+bh*.22,bw*.48,bh*.63,'BFE8F5','5F97AF')
        shape(s,bx+bw*.12,by+bh*.3,bw*.28,bh*.31,'CCB9EC','8367AD',MSO_SHAPE.OVAL)
        shape(s,bx+bw*.21,by+bh*.4,bw*.07,bh*.08,'9272BC',kind=MSO_SHAPE.OVAL)
        for k in range(2):shape(s,bx+bw*(.28+k*(.04 if plant else .4)),by+bh*(.13+k*(.75 if plant else .65)),bw*.16,bh*.09,'F5CA6C','C3933A',MSO_SHAPE.OVAL)
        for k in range(3):line(s,bx+bw*.17,by+bh*(.64+k*.035),bx+bw*.42,by+bh*(.64+k*.035),'B67AAB',1.6)
        for k in range(3):shape(s,bx+bw*(.2+k*.07),by+bh*.64,.035,.035,'353B48',kind=MSO_SHAPE.OVAL)
        points=[(bx+.3,by+.1) if plant else (bx+.1,by+bh*.5),(bx+bw*.55,by+bh*.94) if plant else (bx+bw*.28,by+bh*.72),(bx+bw*.26,by+bh*.44),(bx+bw*.4,by+bh*.17),(bx+bw*.32,by+bh*.69),(bx+bw*.7,by+bh*.17)]
        names=['Cell membrane','Cytoplasm','Nucleus + nuclear envelope','Mitochondrion','Rough endoplasmic reticulum','Ribosomes']
        if plant:
            for k in range(2):shape(s,bx+bw*.1,by+bh*(.12+k*.63),bw*.12,bh*.1,'5DAB62','348348',MSO_SHAPE.OVAL)
            names+=['Cell wall (cellulose)','Large central vacuole','Chloroplast'];points +=[(bx+.01,by+bh*.8),(bx+bw*.63,by+bh*.48),(bx+bw*.16,by+bh*.17)]
    # A few ribosomes scattered clear of the nucleus and vacuole.
    for k in range(9):
        xx=bx+bw*(.1+(k%3)*.09); yy=by+bh*(.68+(k//3)*.065)
        if not bact:xx=bx+bw*(.57+(k%3)*.06);yy=by+bh*(.12+(k//3)*.045)
        shape(s,xx,yy,.035,.035,'353B48',kind=MSO_SHAPE.OVAL)
    for i,((px,py),name) in enumerate(zip(points,names),1):
        shape(s,px-.08,py-.08,.23,.23,'FFFFFF','606B83',MSO_SHAPE.OVAL);text(s,str(i),px-.08,py-.055,.23,.19,8,bold=True,align=PP_ALIGN.CENTER)
        text(s,f'{i}  {name}',x+w*.64,y+.1+(i-1)*(h-.25)/len(names),w*.35,(h-.3)/len(names),12 if w<7 else 18)
    return names


def build2():
    b=Builder(2)
    b.opening('What is a cell?','Cell theory — and what the cells in our diagrams share.', ['State the three ideas of cell theory.','Compare prokaryotic and eukaryotic cells.','Use cell theory to explain new skin cells.'],'cell: basic unit of life\ncell theory: three connected ideas\norganelle: a cell structure with a job', [('Predict',8,'Write, pair-share, harvest predictions'),('Observe',30,'Kitchen mapping; diagrams and reading; 12-card sort'),('Explain',18,'Revise, colour-code, write a CER'),('Driving Question Board',8,'Star progress; post question OR understanding'),('Model Building',12,'Version 1, key and two limitations'),('Summary',4,'Complete the Lesson 2 row')])
    s=b.slide('Predict','What do cells have in common?',4.5,'Give the full 3-minute silent writing period. Then pairs share for 90 seconds. Use a typical-cell comparison; platelets are fragments, not complete cells.',say='Think back to Amina’s finger. Write three sentences: I think all cells share… because…',wait='3 full minutes of silent writing; 90 seconds of partner talk.',look='Every pupil has three starting sentences. Note membrane, DNA or control-centre ideas; flag blanks for support.',iff='If a pupil writes I do not know, record that starting point and support them during diagram study; do not give the comparison now.',why='Writing first preserves each pupil’s own starting point.',key='This is an uncorrected starting prediction, not a test. Revisit after gathering evidence.',phase_note='PREDICT PHASE (8 minutes total, including opening)',check=True)
    card(s,'Write alone first','“I think all cells share… because…”\n\nWrite three sentences.',.65,2.25,8.0,2.6,'E3F7FC',27);timer(s,'3:00',9.3,2.45);text(s,'Then pair-share: 1:30',9.3,3.35,3.3,1,21)
    s=b.slide('Predict','OUR PREDICTIONS',2,'Cold-call three pairs across the room. Record key words in a board column. Do not correct until evidence is available.',say='What did your pair agree on? Let us hold that idea and test it.',wait='5–10 seconds before taking each answer.',look='Whose claims include a reason? Record recurring ideas without endorsing them.',iff='If a pupil says every cell has a nucleus, write it as a prediction to test against the bacterial diagram.',why='The visible list creates a record pupils can revise later.',key='Record pupils’ own statements; no correct class list yet.',check=True)
    card(s,'Before — our predictions','Tell us your idea and your reason.\nYour teacher will record your words.\nWe will test them today.',.65,2.3,12.05,3.3,'E3F7FC',29)
    s=b.slide('Observe','A kitchen can help us think about cells',8,'Groups of four draw a two-column mapping table. Use the five mappings in the key. Then discuss where the analogy fails. Check two tables per group.',say='A home kitchen in Kericho and a hospital canteen in Nairobi use familiar tools in different arrangements. Which feature controls entry?',wait='10 seconds for reflection; 12 seconds after the membrane question.',look='At least five mappings; pupils distinguish recipe information from the place it is stored.',iff='If pupils give a cell a thinking boss, state: Cell parts respond to chemical signals; a cell has no cook making decisions.',why='A familiar analogy supports first ideas, but its limits prevent literal misconceptions.',key='Door/boundary → cell membrane; recipe information → DNA (held in the nucleus in the eukaryotic example); jiko → mitochondria; pots and utensils → organelles; workspace → cytoplasm. Limits: kitchens have cooks, doors are not selectively permeable membranes, kitchens do not divide.',science='The analogy applies to the eukaryotic example. Bacteria have no mitochondria or nucleus. A jiko is a rough energy-release analogy, not an identical process.',phase_note='OBSERVE PHASE (30 minutes total)',check=True)
    table(s,['Kitchen feature','Your cell mapping'],[['Door / boundary',''],['Recipe information',''],['Jiko (stove)',''],['Sufuria, mwiko and utensils',''],['Workspace','']],h=3.7,widths=[5.0,7.05],size=21);text(s,'Where does the analogy stop working?',.65,6.2,12,.45,23,bold=True)
    for kind,title in [('animal','A generalised animal cell'),('plant','A generalised plant cell'),('bacterium','A bacterium: E. coli')]:
        s=b.slide('Observe',title,2,'Use the labelled drawing as evidence, not a real observation through a microscope. Pupils identify names before organelle functions are taught in Lessons 4–5.',say='Look for boundaries, the location of DNA and structures shared with the other drawings.',wait='Allow a quiet scan before inviting observations.',look='Pupils use the actual labels and distinguish a nucleus from bacterial DNA.',iff='If ribosomes are mistaken for membrane-bound compartments, point to the small dots and return to this during Explain.',key='Typical animal: membrane, cytoplasm, nucleus, mitochondria, ER and ribosomes. Typical photosynthetic plant example also shows a cellulose wall, chloroplasts and large central vacuole. The bacterium has a peptidoglycan wall, membrane, cytoplasm, ribosomes and circular chromosome in a nucleoid, with no nuclear membrane.',science='Representative drawings, not to scale. Not every specialised cell has every illustrated structure. Some E. coli strains cause illness, while many are harmless. Do not infer that all bacteria are pathogens.')
        diagram(s,kind,.8,2.05,11.7,4.25);text(s,'Editable illustrative diagram · not a micrograph · not to scale',.65,6.5,12,.3,12,'697080')
    s=b.slide('Observe','Reading: three contributions to cell theory',2,'Read the short information cards before filling the evidence table on the next slide. Provide printed copies for reference while pupils write. These cards supply offline reading.',say='What did each scientist contribute, and how could the idea help us think about Amina’s finger?',wait='Quiet reading, then move to the table.',look='Pupils locate a contribution for each scientist rather than copying a name only.',iff='If reading is difficult, read each card aloud and let pupils point to the matching scientist.',key='Schleiden: plant tissues are built from cells. Schwann: animal tissues are built from cells. Virchow: new cells arise from existing cells.',science='The dates and attributions follow the supplied example: Schleiden 1838, Schwann 1839, Virchow 1855. These are selected historical contributions, not a claim that three people alone established cell theory.')
    for i,(title,body) in enumerate([('Schleiden · Germany · 1838','Studied plants. Plant tissues are built from cells.'),('Schwann · Germany · 1839','Studied animals. Animal tissues are built from cells too.'),('Virchow · Germany · 1855','New cells arise from cells that already exist.')]):card(s,title,body,.65+i*4.1,2.25,3.85,3.7,size=26)
    s=b.slide('Observe','Cell Theory Evidence Table',6,'Pupils complete each scientist row using the reading. The phenomenon-connection sentence is essential. Circulate without filling cells for them.',say='For each scientist, write how the idea connects to Amina’s healing finger.',wait='Wait 10 seconds after individual prompts.',look='Every row includes a contribution and a meaningful connection; read one table per group.',iff='For a blank connection ask the source prompt: Virchow said all cells come from pre-existing cells — how does that explain where new skin cells came from?',why='The last column connects historical knowledge to an explanation.',key='Schleiden, Germany/1838: plants are cellular; the onion Amina sliced consists of plant cells. Schwann, Germany/1839: animals are cellular; Amina’s skin is made of animal cells. Virchow, Germany/1855: new cells come from existing cells; existing skin cells divide to produce new cells during healing. These are source-derived connections, not direct evidence of division in the drawing.',check=True)
    table(s,['Scientist','Country / period','Contribution','Connection to Amina'],[['Schleiden','','',''],['Schwann','','',''],['Virchow','','','']],h=3.65,widths=[2.0,2.1,3.45,4.5],size=19);text(s,'Use the reading. The last column matters most.',.65,6.2,10,.45,23);timer(s,'6:00')
    cards=['Has a membrane-bound nucleus','Main chromosome in a nucleoid','Has mitochondria','Bacterial cell','Has ribosomes','Has a cell membrane','Has membrane-bound organelles','Typical size: 1–5 μm','Typical size: 10–100 μm','Plant, animal or fungal cell','Main chromosome is circular','Has cytoplasm']
    s=b.slide('Observe','Sort 12 features — justify each choice',8,'Give 12 pre-cut cards per group. Sort into Prokaryote / Eukaryote. Put shared cards across the dividing line or copy them into both columns. Give 8 minutes; note disputes for Explain.',say='Use the diagrams. A feature can belong in both categories. Record your final arrangement.',wait='8 minutes of group sorting.',look='Which groups put ribosomes only in eukaryotes? Are membrane and cytoplasm shared?',iff='Do not announce disputed answers yet. Ask: Which diagram gives you evidence for that choice?',key='Prokaryote example: nucleoid, bacterial cell, 1–5 μm, circular main chromosome. Eukaryote examples: nucleus, mitochondria, membrane-bound organelles, 10–100 μm, plant/animal/fungal cell. Both: ribosomes, cell membrane, cytoplasm. The size ranges are typical examples, not absolute classification tests.',check=True)
    for i,c in enumerate(cards):
        x=.65+(i%4)*3.08;y=2.12+(i//4)*1.13;shape(s,x,y,2.87,.96,'F5F6F9','D7DCE5');text(s,c,x+.12,y+.17,2.63,.71,18,bold=True,align=PP_ALIGN.CENTER)
    text(s,'PROKARYOTE  |  EUKARYOTE · Shared cards span the line.',.65,5.75,10,.5,20);timer(s,'8:00',10.7,6.15)
    s=b.slide('Explain','Shared does not mean identical',3,'Revisit the predictions and disputed sort cards. Pupils mark confirmed/revised ideas in a new colour. Correct the actual misconception observed, not an assumed class error.',say='Ribosomes appear in both diagrams. That does not make the two cell types identical.',wait='5–10 seconds after questions.',look='Pupils distinguish ribosomes from membrane-bound organelles and revise their own prior claims.',iff='If ribosomes remain eukaryote-only, point to the bacterial dots and ask the pupil to update the card placement.',key='Both examples show membrane, cytoplasm and ribosomes. The bacterium has no membrane-bound nucleus; the eukaryotic examples do. Only the plant and bacterial examples have a wall, made from different materials.',science='Ribosomes make proteins; their presence alone does not establish pathogenicity. The draft removes that implication from the example.',phase_note='EXPLAIN PHASE (18 minutes total)',check=True)
    table(s,['Feature','Prokaryotic example','Eukaryotic examples'],[['Membrane, cytoplasm, ribosomes','Present','Present'],['Membrane-bound nucleus','Absent','Present'],['Membrane-bound organelles','Absent','Present'],['Cell wall','E. coli: peptidoglycan','Plant: cellulose; animal: none']],h=3.8,widths=[3.65,4.1,4.3],size=20)
    s=b.slide('Explain','Cell theory: three connected ideas',3,'Consolidate the three source tenets. Connect the third to new skin cells; do not explain specialised healing roles yet.',say='Which idea explains where new skin cells came from?',wait='10 seconds before answering.',look='Pupils connect existing cells with new cells rather than saying cells appear from nothing.',iff='If pupils propose spontaneous appearance, return to the third tenet and the observed repair story.',key='All organisms consist of one or more cells. A cell is the basic unit of living structure and function. New cells arise from existing cells. The third, associated here with Virchow, explains that new skin cells derive from existing cells.')
    for i,t in enumerate(['Living things consist of one or more cells.','The cell is the basic unit of structure and function.','New cells arise from cells that already exist.']):
        shape(s,.65,2.1+i*1.25,.6,.6,'CF128C',kind=MSO_SHAPE.OVAL);text(s,str(i+1),.65,2.2+i*1.25,.6,.4,22,'FFFFFF',True,PP_ALIGN.CENTER);text(s,t,1.65,2.17+i*1.25,10.7,.9,27,bold=True)
    s=b.slide('Explain','Colour-code your three drawings',4,'Pupils annotate their diagram sheets, not the reference slide. Keep the source colour system consistent. Explain the cell-wall exception and distinguish the two materials.',say='Green means shared across the two broad cell categories; it does not mean every cell in every category has that structure.',wait='Give pupils time to check label by label with a partner.',look='Ribosomes green; membrane and cytoplasm green. Plant and bacterial walls green, with materials noted; no animal wall. Chloroplasts and central vacuole red.',iff='If the wall is called plant-only, compare the E. coli and plant labels. If nucleus is green, revisit the nucleoid.',key='GREEN: membrane, cytoplasm, ribosomes, and walls on the plant/bacterial examples only. BLUE: nucleus, mitochondria and ER in the eukaryotic examples. RED: chloroplasts and the large central vacuole on the plant example. Plant walls are cellulose; the E. coli wall includes peptidoglycan.',science='The source uses category-level sharing; this requires an explicit exception for walls. Generic diagrams must not be presented as universal descriptions of specialised cells.',check=True)
    for i,(title,body,c) in enumerate([('GREEN','Shared across categories','299354'),('BLUE','Eukaryotic examples only','347BC2'),('RED','Plant example only','DC4545')]):shape(s,.65,2.15+i*1.15,2.3,.6,c);text(s,title,.85,2.24+i*1.15,1.9,.4,20,'FFFFFF',True);text(s,body,3.3,2.18+i*1.15,8.9,.8,26)
    text(s,'Wall exception: plant + bacterium have walls; the animal cell does not.',.65,6.05,12,.55,21,bold=True)
    s=b.slide('Explain','How did Amina’s body make new skin cells?',8,'Give 6 full minutes of silent writing, then ask two pupils to read their work. Collect six notebooks, two from each ability group, for review.',say='Write four sentences. Make a claim, name the observed change, then use cell theory to explain it.',wait='6 minutes silent writing. After readings, allow a brief class thumbs response on the reasoning.',look='A clear claim; observation from the illustrated story as evidence; cell theory used in reasoning with because/therefore.',iff='If a pupil writes only Virchow as the evidence, ask: What change did the story show? Then ask what scientific idea explains it.',why='An observation and an explanatory theory play different roles.',key='Claim: new skin cells came from existing skin cells. Evidence: the story shows new skin appearing as the cut closes. Reasoning: because new cells arise from existing cells, existing skin cells divided to make new ones. Therefore new cells helped close the surface. The drawing illustrates appearance; it does not directly show division or migration.',science='No photograph or observed edge-to-centre spreading is claimed: L1 used illustrative finger drawings. This is a careful limitation of the supplied CER exemplar.',check=True)
    for i,(t,body) in enumerate([('CLAIM','Where did the new skin cells come from?'),('EVIDENCE','What change did the healing story show?'),('REASONING','How does cell theory connect that change to your claim?')]):card(s,t,body,.65+i*4.1,2.2,3.85,3.2,'FBE7F4',24)
    text(s,'Four sentences · use “because” or “therefore”.',.65,5.8,9.5,.5,23,bold=True);timer(s,'6:00')
    s=b.slide('Driving Question Board','What can we now begin to answer?',8,phase_note='DRIVING QUESTION BOARD PHASE (8 minutes total)',do='Reread the driving question. Give 3 minutes silent reflection. Each pupil stars an answered/partly answered question and posts ONE note: question OR growing understanding. Read three or four notes and photograph the board.',say='What has today’s evidence helped us understand? What still puzzles you?',wait='3 minutes of individual reflection.',look='Every pupil posts at least one note. Check green notes for the mistaken claim that all cells are identical.',iff='If a note says all cells are the same, ask the pupil to name a difference visible in today’s diagrams and revise the note.',key='Valid understanding: new cells come from existing cells; typical bacterial and body cells differ in nucleus/organelles. Valid wonder: a genuine unresolved question. Do not claim the full specialisation question is answered.',science='The source asks one new note of either colour, not one of each. An additional evidence vote is not in the source and is flagged in the build log.',check=True)
    card(s,'1 · Star progress','Which earlier question has today begun to answer?',.65,2.2,3.85,3.4,'FCBA63',24);card(s,'2 · Choose ONE note','Yellow: a new question\nOR\nGreen: growing understanding',4.75,2.2,3.85,3.4,'FFF1A3',23);card(s,'3 · Keep investigating','What do we still need to explain about Amina’s finger?',8.85,2.2,3.85,3.4,'D8EDC9',24);timer(s,'3:00')
    s=b.slide('Model Building','Version 1: a model that can grow',12,phase_note='MODEL BUILDING PHASE (12 minutes total)',do='Pupils use the same notebook record, dated Version 1. Give 10 minutes independent sketch/key work and 2 minutes partner limitation sharing. Collect all models for review; return next lesson.',say='The required labels are membrane, nucleus, cytoplasm, and cell wall if you choose a plant cell. Other structures are optional at this stage.',wait='10 minutes independent work; wait 10 seconds after a correction prompt.',look='Three required animal labels or four plant labels; role statements; at least two genuine limitations. Functions beyond the core are provisional ideas, revisited in L4–5.',iff='If an animal model has a wall, ask: What did the animal diagram show? If the key only states a location, ask for the structure’s possible role.',key='Membrane: boundary controlling movement in/out; nucleus: holds DNA; cytoplasm: internal material where many reactions occur; plant wall: rigid outer support. Optional structures need not be fully understood yet. Limitations are individual honest gaps, not a fixed answer.',science='The V2 source requires 3–4 core labels, unlike the example’s longer compulsory list. Detailed functions are taught later. Language support from V2: confirm proposed Kiswahili terms with the school’s language teacher before use. Pupils may discuss in Kiswahili or a home language first, then English; accept a labelled drawing or spoken explanation where appropriate.',check=True)
    card(s,'Draw + label','Membrane\nNucleus\nCytoplasm\nWall if plant\nExtras optional',.65,2.2,3.85,3.6,'FCE8EA',23);card(s,'Write a model key','One sentence per label: your best idea of its role.\n\nThese ideas will be revised.',4.75,2.2,3.85,3.6,'FCE8EA',23);card(s,'Name two limits','“This model does not yet show…”\n\nShare your limits with a partner.',8.85,2.2,3.85,3.6,'FCE8EA',23);timer(s,'10:00')
    summary(b,4,'Next lesson: compare E. coli with a human cheek cell.')
    b.save(80)


def summary(b,minutes,nextline):
    l=b.lesson;key='\n'.join(k+': '+v for k,v in l['summaryTablePrompt'].items())
    if b.n==3:
        key='Observed: labelled E. coli and cheek-cell diagrams; nucleus/envelope, organelles, wall and typical size compared.\nLearned: the bacterial chromosome is in a nucleoid with no nuclear membrane; the cheek-cell nucleus is membrane-bound. Both examples have ribosomes, membrane and cytoplasm.\nExplained: bacteria that could enter a cut and typical human body cells have different cellular organisation. This comparison alone does not prove how immune recognition or wound repair works; those mechanisms remain for later lessons.'
    s=b.slide('Driving Question Board',f'Summary Table · Lesson {b.n}',minutes,'Pupils complete the lesson row. Bring the explanation back to Amina’s cut, then close with the next lesson teaser.',say='What did we observe? What did we learn? What can we now explain?',look='Answers refer to this lesson’s evidence and distinguish observations from ideas.',iff='If the phenomenon link is absent, ask which comparison or theory idea helps explain Amina’s healing finger.',key=key,reflection=l['teacherReflection'],dark=True)
    for i,(title,body) in enumerate([('What did I observe?','Name the diagrams, patterns or activity evidence.'),('What did I learn?','Which idea or comparison can I now explain?'),('How does this explain it?','Connect the evidence to Amina’s finger.')]):card(s,title,body,.65+i*4.1,2.25,3.85,3.3,'292929',22)
    # Text inside dark cards must be light.
    for sh in list(s.shapes):
        if sh.has_text_frame and sh.top>=Inches(2.25) and sh.top<Inches(5.6):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:r.font.color.rgb=rgb('FFFFFF')
    text(s,nextline,.65,6.0,12,.65,22,'FCBA63',True)


def build3():
    b=Builder(3)
    b.opening('Two cell types','E. coli and a human cheek cell — compare the structures, then explain why they matter.', ['Compare a bacterium with a human cheek cell.','Use structural evidence in a comparison table.','Connect the comparison to Amina’s finger.'],'prokaryote: lacks a membrane-bound nucleus\neukaryote: has a membrane-bound nucleus\nnucleoid: bacterial DNA region', [('Predict',5,'Write two reasons; share and harvest'),('Observe',10,'Compare diagrams; complete and check table'),('Explain',10,'Five cold-calls; three sentences; terms'),('Driving Question Board',5,'What we know / still wonder'),('Model Building',6,'Version 1 add-on, four labels per cell'),('Summary',4,'Complete the Lesson 3 row')])
    s=b.slide('Predict','Are bacteria built like Amina’s body cells?',3,'Give 2 minutes silent writing, then 1 minute shoulder-partner sharing. Predictions must include two reasons. Do not correct yet.',say='Amina wonders about bacteria entering her cut. Predict whether they are built the same way as her skin cells.',wait='2 full minutes silent writing, then 1 minute partner talk.',look='Use of nucleus, membrane and organelle vocabulary from Lesson 2; two reasons written.',iff='If a pupil uses only everyday words, keep that prediction and scaffold with diagram labels during Observe.',key='No single correct prediction is supplied before evidence; return to and revise both reasons during Explain.',phase_note='PREDICT PHASE (5 minutes total including opening)',check=True)
    card(s,'Write your prediction','Same or different?\nGive TWO reasons.\nThen share with a partner.',.65,2.25,8.5,3.3,'E3F7FC',29);timer(s,'2:00',10.0,2.5);text(s,'Partner talk: 1:00',9.5,3.4,3.2,.7,22)
    s=b.slide('Predict','Before — our predictions',1,'Use the class register to call two pupils. Scribe key words without correction. Tell the class these will be tested today.',say='Which reason did you and your partner discuss?',look='Two reasoned ideas recorded on the board.',iff='If only size is mentioned, keep it as a claim to test rather than naming a new feature for them.',key='Pupil-generated predictions; not a corrected comparison yet.')
    text(s,'Share one prediction and your reason.\nWe will test it using the drawings.',.8,2.7,11.7,2.0,34,bold=True)
    s=b.slide('Observe','Compare: E. coli and a cheek cell',3,'Use printed copies of this slide. Model the membrane-bound-nucleus row only, then pupils read the remaining labels. No live bacterial culture or cheek sampling occurs.',say='E. coli is a bacterium. Some strains can cause illness; many are harmless. This is our bacterial example, not a diagnosis of Amina’s cut.',wait='At least 5 seconds after a diagram question.',look='Nucleoid is not labelled as a membrane-bound nucleus. Ribosomes are present in both.',iff='If pupils confuse the nucleoid and nucleus, trace the actual nuclear envelope on the cheek-cell drawing and note its absence around bacterial DNA.',key='Bacterium: nucleoid, no membrane-bound nucleus/organelles, peptidoglycan wall. Cheek cell: membrane-bound nucleus and organelles, no cell wall. Both have ribosomes, membrane and cytoplasm.',science='Representative diagram labels include structures too small to resolve in a school light microscope. L7 is the actual cheek-cell observation. No fixed 50-fold size ratio is inferred from these drawings.',phase_note='OBSERVE PHASE (10 minutes total)')
    text(s,'E. coli · prokaryote',.8,1.97,5.8,.35,20,bold=True);text(s,'Human cheek cell · eukaryote',6.9,1.97,5.8,.35,20,bold=True)
    diagram(s,'bacterium',.75,2.45,5.85,3.6);diagram(s,'animal',6.8,2.45,5.85,3.6)
    text(s,'E. coli: approximately 1–5 μm\nCheek-cell example: approximately 10–100 μm · drawings not to scale',.75,6.22,12,.55,15,'697080')
    features=['Membrane-bound nucleus','Membrane-bound organelles','Approximate size','Cell wall + material','Location of main DNA']
    s=b.slide('Observe','Use the diagrams to complete your table',5,'Pupils work independently for 4 minutes, then compare for 1 minute. Place a question mark beside disagreements. Do not correct during independent work.',say='Use the diagram labels as evidence for each row.',wait='4 minutes individual, then 1 minute partner comparison.',look='Five rows complete, with nucleoid distinct from nucleus. Also ask whether ribosomes occur in both.',iff='If a row is disputed, mark it for the review rather than guessing. Read diagrams, not the answer slide.',key='No membrane-bound nucleus / yes; no membrane-bound organelles / yes; approximately 1–5 μm / 10–100 μm in these examples; peptidoglycan wall / no wall; main chromosome in nucleoid / main chromosomes inside nucleus.',check=True)
    table(s,['Feature','E. coli','Human cheek cell'],[[f,'',''] for f in features],h=3.8,widths=[4.6,3.75,3.7],size=20);text(s,'Work alone: 4:00 · compare: 1:00 · mark disagreements with ?',.65,6.15,12,.5,21)
    s=b.slide('Observe','Check the disputed rows',2,'Review rows with question marks and always check the ribosome misconception. Call across seating groups, not only volunteers.',say='Which label supports your answer?',wait='5 seconds before prompting.',look='Pupils justify an answer with an actual diagram feature.',iff='If a pupil says no nucleus means no DNA, point to the chromosome label. If ribosomes are called membrane-bound, distinguish the dots from compartments.',key='Nucleoid: DNA region without a nuclear envelope. Ribosomes: both, not membrane-bound. The bacterial wall includes peptidoglycan; no wall surrounds the cheek cell.',check=True)
    card(s,'DNA ≠ nucleus','A bacterium has DNA.\nIts DNA is not enclosed in a nucleus.',.65,2.3,5.85,3.3,'F4F5F8',28);card(s,'Ribosomes: BOTH','Ribosomes do not have an enclosing membrane.',6.85,2.3,5.85,3.3,'F4F5F8',28)
    questions=[('Where is E. coli’s DNA?', 'How does that region differ from a cheek-cell nucleus?',.8,'5 seconds.','The nucleoid contains the main bacterial chromosome and is not surrounded by a nuclear envelope. The cheek-cell nucleus has an envelope.'),('Both cells have ribosomes.','Does that make them the same cell type? Explain.',1.2,'7 seconds.','No. Ribosomes are shared structures; the nucleus and membrane-bound organelles distinguish the examples.'),('Choose one structural difference.','Use your table to explain the difference.',1,'5 seconds.','Accept a supported row from the comparison table: nucleus, organelles, wall/material, typical size or DNA location. Shared ribosomes are not a difference.'),('Why might a larger cell benefit from compartments?','Use the drawings to suggest a reason.',1.5,'8 seconds.','Membrane-bound compartments separate different processes. Treat this as a reasoning question; a size difference alone does not classify the cell.'),('The bacteria in our example are which type?','How does the diagram support your answer?',1.5,'5 seconds.','E. coli is prokaryotic: its DNA is in a nucleoid without a nuclear membrane and it lacks membrane-bound organelles. Amina’s typical nucleated body cells are eukaryotic.')]
    for i,(title,prompt,t,wait,key) in enumerate(questions):
        s=b.slide('Explain',f'Question {i+1} of 5',t,'Cold-call a different pupil for this question using the register. After the response ask everyone for thumbs up/sideways/down; ask one pupil to justify a differing signal. '+('Return to the written predictions and revise them in a new colour.' if i==0 else ''),say=title+' '+prompt,wait=wait,look='An explanation based on a labelled structure, not only a recalled category word.',iff='If uncertain, point to the relevant label and ask the pupil to read the evidence. Revisit the shared ribosome example if needed.',key=key,science='The five source questions are retained. The fourth omits the unsupported exact 50-times size claim. Do not present prokaryotes as failed or undeveloped eukaryotes.',phase_note='EXPLAIN PHASE (10 minutes total)' if i==0 else '',check=True)
        text(s,title,.8,2.3,11.7,1.55,36,bold=True);text(s,prompt,.8,4.3,11.6,1.45,29);text(s,'Everyone: agree, partly agree or disagree — and why?',.8,6.1,11.7,.6,20,'697080')
    s=b.slide('Explain','Write three connected sentences',3,'Give 3 full minutes independent writing. Read over shoulders; identify pupils whose third sentence is only definitional for targeted model questioning.',say='Connect the cell comparison to Amina’s healing-cut story.',wait='3 minutes silent writing.',look='Sentence 3 links bacteria and body cells to the phenomenon using structural evidence.',iff='If all sentences are definitions, ask: Which cells belong to Amina, and which could enter from outside?',key='E. coli has no membrane-bound nucleus, making it a prokaryote. The cheek cell has a membrane-bound nucleus, making it a eukaryote. This helps us distinguish bacterial cells that could enter the cut from Amina’s own body cells; they are both living cells but have different organisation.',check=True)
    for i,t in enumerate(['One feature that makes E. coli a prokaryote.','One feature that makes the cheek cell a eukaryote.','Why this comparison matters for Amina’s finger.']):text(s,f'{i+1}.  '+t,.8,2.3+i*1.1,11.7,.85,28,bold=True)
    timer(s,'3:00')
    s=b.slide('Explain','Two names, one memory aid',1,'Introduce the source word roots after the evidence and written explanation. Pupils copy them as a memory anchor.',say='Prokaryote: before + nucleus. Eukaryote: true + nucleus. Use the structures, not the word roots alone, to identify a cell.',key='Pro = before; eu = true; karyon = nucleus. The modern distinction here is presence or absence of a membrane-bound nucleus.',science='The etymology is a mnemonic, not a claim that living bacteria are incomplete or inferior cells.')
    card(s,'PROKARYOTE','pro = before\nkaryon = nucleus\n\nNo membrane-bound nucleus.',.65,2.25,5.85,3.7,'E3F7FC',27);card(s,'EUKARYOTE','eu = true\nkaryon = nucleus\n\nMembrane-bound nucleus.',6.85,2.25,5.85,3.7,'FBE7F4',27)
    s=b.slide('Driving Question Board','What we know — and still wonder',5,phase_note='DRIVING QUESTION BOARD PHASE (5 minutes total)',do='Read two or three earlier relevant wonder notes. Pupils write the two source stems for 2 minutes, post them in KNOW and WONDER columns, then spend 2 minutes scanning and reading three wonder notes.',say='Be specific: write HOW the cells differ, using today’s structural vocabulary.',wait='2 minutes writing; 2 minutes gallery scan.',look='KNOW statements name nucleus, organelles or wall rather than just bigger or more complex.',iff='If a note says only different, ask: Which feature in your table supports that?',key='KNOW: the typical bacterial and cheek-cell examples differ in the nuclear envelope and membrane-bound compartments. WONDER: individual questions about functions/structures that remain unresolved. Do not answer L4’s organelle questions now.',check=True)
    card(s,'WHAT WE NOW KNOW','“Prokaryotes and eukaryotes differ in…”',.65,2.25,5.85,3.3,'FFF1A3',28);card(s,'WHAT WE STILL WONDER','“But we still wonder…”',6.85,2.25,5.85,3.3,'DFF4FA',28);timer(s,'2:00')
    s=b.slide('Model Building','Version 1 add-on: two cell types',6,phase_note='MODEL BUILDING PHASE (6 minutes total)',do='Add a second panel beside/below the existing model. Give 5 minutes drawing, then select two pupils for a brief shared annotation. Keep four labels per cell and the caption.',say='This is an addition to your current model, not a replacement. Use the table as your evidence.',wait='5 minutes drawing; 1 minute sharing.',look='At least four labels per cell. Ribosomes green. Caption explains a consequence, not merely that the categories differ.',iff='If the bacterial DNA is missing, ask: No nucleus — or no DNA? If a wall appears on both E. coli and cheek cell, check the table. Plant and bacterial walls, when compared, have different materials.',key='Orange: bacterial-specific features in this pair, such as the nucleoid and peptidoglycan wall. Blue: cheek-cell nucleus and membrane-bound organelles. Green: membrane, cytoplasm, ribosomes. Caption: a nuclear envelope and internal compartments organise eukaryotic processes differently. No correct artistic outline is required.',science='Source colour scheme is pair-specific. Do not colour the bacterial wall as unique to all prokaryotes versus all eukaryotes: plants also have walls.',check=True)
    card(s,'Draw + label','E. coli and a cheek cell\nAt least FOUR labels each.\nUse your comparison table.',.65,2.2,3.85,3.5,'FCE8EA',23);card(s,'Colour the comparison','ORANGE: bacterium only\nBLUE: cheek cell only\nGREEN: both',4.75,2.2,3.85,3.5,'FCE8EA',22);card(s,'Write a caption','“A prokaryote has no membrane-bound nucleus, while a eukaryote does — this matters because…”',8.85,2.2,3.85,3.5,'FCE8EA',22);timer(s,'5:00')
    summary(b,4,'Next lesson: inside a plant cell — return to Amina’s onion.')
    b.save(40)

if __name__=='__main__':
    build2()
    build3()
